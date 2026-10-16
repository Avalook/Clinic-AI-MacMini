"""Thống kê THEO HÀNG HOÁ cho kế toán (Tuyền 14/10/2026, Q18).

Bảng tổng theo mặt hàng trong kỳ: mã, tên, ĐVT, số lượng, tiền trước giảm, giảm
giá, tiền sau giảm (+ cột thuế để G5b điền). Nguồn là CHÍNH các dòng hoá đơn mà
màn Xuất hoá đơn dùng (`XuatHoaDonService`): cùng phạm vi quyền, cùng luật bỏ
phiếu huỷ / trừ dòng đã hoàn — hai màn không thể lệch nhau.
"""

from __future__ import annotations

import io
import re
from collections.abc import Mapping, Sequence
from typing import Any

import asyncpg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from clinicai.api.identity import StaffIdentity
from clinicai.services.thue_gtgt import chuan_hoa_thue_suat, tach_thue
from clinicai.services.xuat_hoa_don_service import (
    NHOM_DICH_VU,
    NHOM_THUOC,
    NHOM_VAT_TU,
    XuatHoaDonService,
)

#: Thứ tự nhóm trên màn / file.
NHOM_THU_TU = (NHOM_DICH_VU, NHOM_THUOC, NHOM_VAT_TU)
TEN_NHOM = {
    NHOM_DICH_VU: "Dịch vụ",
    NHOM_THUOC: "Thuốc",
    NHOM_VAT_TU: "Vật tư",
}
_KHOANG_TRANG = re.compile(r"\s+")


def _khoa_ten(ten: str) -> str:
    """Tên so khớp: bỏ khoảng trắng thừa, không phân biệt hoa thường."""
    return _KHOANG_TRANG.sub(" ", ten.strip()).casefold()


def gom_theo_hang_hoa(
    hoa_don: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    """Dòng của các hoá đơn → bảng theo mặt hàng. Hàm thuần.

    Một mặt hàng = (nhóm, mã nếu có — không thì tên đã chuẩn hoá, ĐVT). Cùng mã
    mà tên khác nhau (đổi tên giữa kỳ) gộp một dòng, giữ tên gặp ĐẦU TIÊN.
    Thuế (G5b): `thue_suat` = thuế suất chung của mặt hàng ("Nhiều" nếu các dòng
    khác nhau, None nếu chưa đặt); `tien_thue` = thuế tách ngược trên tổng SAU giảm
    của từng thuế suất (dòng chưa đặt không có thuế). Thuế KHÔNG đổi các cột tiền.
    """
    nhom: dict[tuple[str, str, str], dict[str, Any]] = {}
    gom_thue: dict[tuple[str, str, str], dict[str | None, int]] = {}
    chua_thue = 0
    for hd in hoa_don:
        for d in hd.get("dong") or []:
            ma = str(d.get("ma") or "")
            ten = str(d.get("ten") or "")
            dvt = str(d.get("dvt") or "")
            loai = str(d.get("nhom") or NHOM_DICH_VU)
            khoa = (loai, ma or _khoa_ten(ten), dvt)
            hang = nhom.setdefault(
                khoa,
                {
                    "nhom": loai,
                    "ma": ma,
                    "ten": ten,
                    "dvt": dvt,
                    "so_luong": 0.0,
                    "tien_truoc_giam": 0,
                    "giam_gia": 0,
                    "tien_sau_giam": 0,
                    "so_dong": 0,
                    "thue_suat": None,
                    "tien_thue": None,
                },
            )
            hang["so_luong"] += float(d.get("so_luong") or 0)
            hang["tien_truoc_giam"] += int(d.get("truoc_giam") or 0)
            hang["giam_gia"] += int(d.get("giam_gia") or 0)
            hang["tien_sau_giam"] += int(d.get("thanh_tien") or 0)
            hang["so_dong"] += 1
            ma_thue = chuan_hoa_thue_suat(d.get("thue_suat"))
            chua_thue += ma_thue is None
            theo_thue = gom_thue.setdefault(khoa, {})
            theo_thue[ma_thue] = theo_thue.get(ma_thue, 0) + int(
                d.get("thanh_tien") or 0
            )
    for khoa, hang in nhom.items():
        theo_thue = gom_thue[khoa]
        co = [m for m in theo_thue if m is not None]
        if co:
            hang["thue_suat"] = (
                co[0] if len(co) == 1 and None not in theo_thue else "Nhiều"
            )
            hang["tien_thue"] = sum(tach_thue(theo_thue[m], m)[1] for m in co)
    thu_tu = {n: i for i, n in enumerate(NHOM_THU_TU)}
    ds = sorted(
        nhom.values(),
        key=lambda h: (thu_tu.get(h["nhom"], 99), h["ten"].casefold(), h["dvt"]),
    )
    for h in ds:
        sl = h["so_luong"]
        h["so_luong"] = int(sl) if float(sl).is_integer() else round(sl, 4)
    return {
        "hang_hoa": ds,
        "so_mat_hang": len(ds),
        "tong_truoc_giam": sum(h["tien_truoc_giam"] for h in ds),
        "tong_giam_gia": sum(h["giam_gia"] for h in ds),
        "tong_sau_giam": sum(h["tien_sau_giam"] for h in ds),
        "tong_thue": sum(int(h["tien_thue"] or 0) for h in ds),
        "dong_chua_thue": chua_thue,
    }


class TheoHangHoaService:
    def __init__(self, pool: asyncpg.Pool) -> None:
        self.pool = pool

    async def xem(
        self,
        *,
        identity: StaffIdentity,
        tu: str | None = None,
        den: str | None = None,
        loai: str | None = None,
        co_so: str | None = None,
    ) -> dict[str, Any]:
        """Cùng cổng quyền / phạm vi với Xuất hoá đơn (không tự dựng luật)."""
        data, goc = await XuatHoaDonService(self.pool).doc_goc(
            identity=identity, tu=tu, den=den, loai=loai, co_so=co_so
        )
        return {
            "tu": data["tu"],
            "den": data["den"],
            "co_so": data.get("co_so"),
            "so_hoa_don": len(goc),
            **gom_theo_hang_hoa(goc),
        }


_TIEU_DE = (
    "Nhóm",
    "Mã hàng",
    "Tên hàng hoá / dịch vụ",
    "ĐVT",
    "Số lượng",
    "Tiền trước giảm",
    "Giảm giá",
    "Tiền sau giảm",
    "Thuế suất",
    "Tiền thuế",
)


def xlsx_theo_hang_hoa(data: Mapping[str, Any]) -> bytes:
    """Excel sửa được: dòng 1 tiêu đề, mỗi mặt hàng một dòng, dòng cuối tổng."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Theo hàng hoá"
    ws.append(list(_TIEU_DE))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for h in data["hang_hoa"]:
        ws.append(
            [
                TEN_NHOM.get(h["nhom"], h["nhom"]),
                h["ma"],
                h["ten"],
                h["dvt"],
                h["so_luong"],
                h["tien_truoc_giam"],
                h["giam_gia"],
                h["tien_sau_giam"],
                h["thue_suat"] or "",
                h["tien_thue"] if h["tien_thue"] is not None else "",
            ]
        )
    ws.append(
        [
            "Tổng",
            "",
            "",
            "",
            "",
            data["tong_truoc_giam"],
            data["tong_giam_gia"],
            data["tong_sau_giam"],
            "",
            data.get("tong_thue") or "",
        ]
    )
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for cot, rong in zip("ABCDEFGHIJ", (10, 14, 44, 8, 10, 16, 14, 16, 10, 14)):
        ws.column_dimensions[cot].width = rong
    # Chữ nhập từ dữ liệu khách (tên hàng) không bao giờ thành công thức.
    for row in ws.iter_rows(min_row=2, max_col=4):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
