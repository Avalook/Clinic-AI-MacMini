"""Editable XLSX tables shared by every data download."""

from collections.abc import Iterable, Mapping, Sequence
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
Sheet = tuple[str, Sequence[str], Iterable[Sequence[Any]]]
_TZ = ZoneInfo("Asia/Ho_Chi_Minh")


def excel_table(sheets: Sequence[Sheet]) -> bytes:
    """Build unlocked tables, preserving numbers, dates and literal text."""
    book = Workbook()
    book.remove(book.active)
    for index, (name, headers, rows) in enumerate(sheets, 1):
        ws = book.create_sheet(name)
        ws.append(list(headers))
        for row in rows:
            ws.append([_cell_value(value) for value in row])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                # Explicit string type prevents Excel formula injection without
                # changing the original patient/name/code text.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "dd/mm/yyyy hh:mm:ss"
                elif isinstance(cell.value, date):
                    cell.number_format = "dd/mm/yyyy"
                elif isinstance(cell.value, (int, float, Decimal)):
                    cell.number_format = "#,##0.########"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        table = Table(displayName=f"DuLieu{index}", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        ws.add_table(table)
        ws.freeze_panes = "A2"
        for column, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(column)].width = min(
                50, max(16, len(header) + 4)
            )
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(_TZ).replace(tzinfo=None)
    return value


def excel_datetime(value: Any) -> datetime | None:
    """Parse an event timestamp without making malformed input fail export."""
    if isinstance(value, datetime):
        return value
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except (ValueError, TypeError):
        return None


def excel_date(value: Any) -> date | None:
    """Parse a report date, returning blank for invalid input."""
    try:
        return date.fromisoformat(str(value))
    except (ValueError, TypeError):
        return None


def xlsx_inventory(report: Mapping[str, Any]) -> bytes:
    """Export stock quantities and values as sortable numeric columns."""
    keys = (
        "ton_dau",
        "nhap",
        "xuat",
        "dieu_chinh",
        "huy",
        "tra_lai",
        "ton_cuoi",
        "gia_tri_ton",
    )
    headers = [
        "Mã hàng",
        "Tên thuốc",
        "Đơn vị",
        "Tồn đầu",
        "Nhập",
        "Xuất bán",
        "Điều chỉnh",
        "Huỷ",
        "Khách trả",
        "Tồn cuối",
        "Giá trị tồn",
    ]
    rows = [
        [
            row.get("ma_hang"),
            row.get("ten"),
            row.get("don_vi") or row.get("don_vi_ban"),
            *[row.get(key) for key in keys],
        ]
        for row in report["dong"]
    ]
    return excel_table([("Xuất nhập tồn", headers, rows)])


def xlsx_theo_doi_thuoc(report: Mapping[str, Any]) -> bytes:
    """Export where a drug is now and its full movement timeline."""
    lo = report["dang_o_dau"]["lo"]
    sheets: list[Sheet] = [
        (
            "Dòng thời gian",
            [
                "Thời gian",
                "Người làm",
                "Hành động",
                "Từ",
                "Đến",
                "Số lô",
                "Hạn dùng",
                "Số lượng",
                "Đơn vị",
                "Tồn lô trước",
                "Tồn lô sau",
                "Tồn nơi đến trước",
                "Tồn nơi đến sau",
                "Khách",
                "Mã BN",
                "Mã phiếu",
                "Lý do",
            ],
            [
                [
                    d["luc"],
                    d["nguoi_lam"],
                    d["hanh_dong"],
                    d["tu"],
                    d["den"],
                    d["so_lo"],
                    d["han_dung"],
                    d["so_luong"],
                    d["don_vi"],
                    d["ton_truoc"],
                    d["ton_sau"],
                    d["ton_den_truoc"],
                    d["ton_den_sau"],
                    d["khach"],
                    d["ma_bn"],
                    d["ma_phieu"],
                    d["ly_do"],
                ]
                for d in report["dong"]
            ],
        )
    ]
    if lo:
        sheets.append(
            (
                "Đang ở đâu",
                [
                    "Số lô",
                    "Hạn dùng",
                    "Đơn vị",
                    "Ở kho",
                    "Ở quầy",
                    "Khách trả chờ kiểm",
                    "Đã huỷ",
                ],
                [
                    [
                        d["so_lo"],
                        d["han_dung"],
                        d["don_vi"],
                        d["o_kho"],
                        d["o_quay"],
                        d["cho_kiem"],
                        d["da_huy"],
                    ]
                    for d in lo
                ],
            )
        )
    return excel_table(sheets)
