"""Xuất hoá đơn mẫu MISA: SQL đọc ĐVT + địa chỉ / CCCD / email trên Postgres thật."""

import io
from dataclasses import replace
from typing import Any
from uuid import uuid4

import asyncpg
import pytest
from openpyxl import load_workbook

from clinicai.api.identity import ClinicRole
from clinicai.core.clock import now_vn
from clinicai.services.quay_thu_service import ma_phieu
from clinicai.services.xuat_hoa_don_service import XuatHoaDonService, xlsx_misa
from tests.services.test_tien_thuoc_cp1_db import tao_quay

pytest_plugins = ["tests.services.test_luot_kham_service_db"]
pytestmark = [pytest.mark.db, pytest.mark.asyncio]


async def test_xuat_hoa_don_doc_dvt_va_thong_tin_khach(pool: asyncpg.Pool) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    cycle, huy = str(uuid4()), str(uuid4())
    for cid, status in ((cycle, "PAID"), (huy, "VOIDED")):
        await pool.execute(
            "INSERT INTO payment_cycle(payment_cycle_id,clinic_id,visit_id,kind,"
            "amount,status,legacy,paid_at,closed_at,created_at,created_by,"
            "confirmed_by,method) VALUES($1::uuid,$2::uuid,$3::uuid,'dich_vu',"
            "350000,$5,true,now(),CASE WHEN $5='VOIDED' THEN now() END,now(),"
            "$4::uuid,$4::uuid,'TRANSFER')",
            cid,
            manager.clinic_id,
            q.visit_id,
            manager.staff_id,
            status,
        )
    for src, ten, unit, gia in (
        ("exam", "Khám phụ khoa", None, 300000),
        ("vat_tu", "Que thử", "Cái", 50000),
    ):
        await pool.execute(
            "INSERT INTO payment_bill_line(clinic_id,payment_cycle_id,visit_id,"
            "kind,source_type,source_id,name_snapshot,quantity,unit,unit_price,"
            "line_total,billing_owner) VALUES($1::uuid,$2::uuid,$3::uuid,"
            "'dich_vu',$4,'x',$5,1,$6,$7,$7,'CLINIC')",
            manager.clinic_id,
            cycle,
            q.visit_id,
            src,
            ten,
            unit,
            gia,
        )
    patient = await pool.fetchval(
        "UPDATE patient p SET address=NULL,address_detail='12 Kim Ngưu',"
        "ward_name='Phường Vĩnh Tuy',province_name='Hà Nội',"
        "national_id_number='001190000001' FROM visit v"
        " WHERE v.visit_id=$1::uuid AND p.clinic_patient_id=v.clinic_patient_id"
        " RETURNING p.clinic_patient_id",
        q.visit_id,
    )
    await pool.execute(
        "INSERT INTO patient_contact_channel(clinic_id,clinic_patient_id,"
        "channel_type,channel_value,is_primary) VALUES($1::uuid,$2,'EMAIL',"
        "'khach@example.com',true)",
        manager.clinic_id,
        patient,
    )
    day = now_vn().date().isoformat()
    data = await XuatHoaDonService(pool).xem(identity=manager, tu=day, den=day)
    ma = {h["ma_phieu"]: h for h in data["hoa_don"]}
    assert ma_phieu(huy) not in ma
    hd = ma[ma_phieu(cycle)]
    assert hd["dia_chi"] == "12 Kim Ngưu, Phường Vĩnh Tuy, Hà Nội"
    assert hd["cccd"] == "001190000001" and hd["email"] == "khach@example.com"
    assert hd["hinh_thuc"] == "Chuyển khoản" and hd["ma_so_thue"] == ""
    assert [(d["ten"], d["dvt"], d["thanh_tien"]) for d in hd["dong"]] == [
        ("Khám phụ khoa", "Lần", 300000),
        ("Que thử", "Cái", 50000),
    ]
    ws: Any = load_workbook(io.BytesIO(xlsx_misa([hd]))).active
    assert ws.cell(10, 7).value == "khach@example.com"
    # Số thứ tự = vị trí trong file (DB chung có phiếu của bài khác) — so theo hd.
    assert [ws.cell(11, c).value for c in (1, 13, 14, 17)] == [
        str(hd["stt"]),
        "Que thử",
        "Cái",
        50000,
    ]
    # Ngày rác → mặc định, không ném.
    assert "hoa_don" in await XuatHoaDonService(pool).xem(
        identity=manager, tu="rác", den="31/02"
    )
