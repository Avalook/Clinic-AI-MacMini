"""Thuế GTGT (G5b): đặt thuế, tách ngược khi xuất, cổng quyền, ràng buộc DB."""

from __future__ import annotations

import io
from dataclasses import replace
from typing import Any
from uuid import uuid4

import asyncpg
import pytest
from openpyxl import load_workbook

from clinicai.api.identity import ClinicRole, StaffIdentity
from clinicai.core.clock import now_vn
from clinicai.core.exceptions import SafetyGateError, ValidationError
from clinicai.services.quay_thu_service import ma_phieu
from clinicai.services.thue_suat_service import ThueSuatService
from clinicai.services.xuat_hoa_don_service import (
    DONG_DAU,
    XuatHoaDonService,
    xlsx_misa,
)
from tests.services.test_quyen_lego_man_db import _nguoi as nguoi_theo_lego
from tests.services.test_tien_thuoc_cp1_db import tao_quay

pytest_plugins = ["tests.services.test_luot_kham_service_db"]
pytestmark = [pytest.mark.db, pytest.mark.asyncio]

GOI_KE_TOAN = ["bao_cao", "ke_toan", "ds_benh_nhan", "lich_lam_viec"]


async def _phieu(
    pool: asyncpg.Pool,
    q: Any,
    manager: StaffIdentity,
    dong: list[tuple[Any, ...]],
    kind: str = "dich_vu",
) -> tuple[str, dict[str, str]]:
    """Một phiếu thu dịch vụ + các dòng (src, tên, ĐVT, SL, đơn giá, id thuốc).
    Trả (id phiếu, {tên: id dòng})."""
    cycle = str(uuid4())
    tong = sum(sl * gia for _, _, _, sl, gia, _ in dong)
    await pool.execute(
        "INSERT INTO payment_cycle(payment_cycle_id,clinic_id,visit_id,kind,"
        "amount,status,legacy,paid_at,created_at,created_by,confirmed_by,method)"
        " VALUES($1::uuid,$2::uuid,$3::uuid,$6,$5,'PAID',true,now(),"
        "now(),$4::uuid,$4::uuid,'CASH')",
        cycle,
        manager.clinic_id,
        q.visit_id,
        manager.staff_id,
        tong,
        kind,
    )
    ids: dict[str, str] = {}
    for src, ten, dvt, sl, gia, th in dong:
        ids[ten] = str(
            await pool.fetchval(
                "INSERT INTO payment_bill_line(clinic_id,payment_cycle_id,visit_id,"
                "kind,source_type,source_id,name_snapshot,quantity,unit,unit_price,"
                "line_total,billing_owner,drug_catalog_id)"
                " VALUES($1::uuid,$2::uuid,$3::uuid,$12,$4,$5,$6,$7,$8,$9,$10,"
                "'CLINIC',$11::uuid) RETURNING id::text",
                manager.clinic_id,
                cycle,
                q.visit_id,
                src,
                f"{src}-{uuid4().hex[:6]}",
                ten,
                sl,
                dvt,
                gia,
                sl * gia,
                th,
                kind,
            )
        )
    return cycle, ids


async def test_dat_thue_nhom_mat_hang_tach_khi_xuat_va_cot_k_l(
    pool: asyncpg.Pool,
) -> None:
    """Chỉ đụng nhóm THUỐC: DB test chung — đặt thuế cho nhóm Dịch vụ / Vật tư sẽ
    đổi số xuất của các bài hoá đơn khác chạy song song."""
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    ke_toan = await nguoi_theo_lego(pool, "CSKH", GOI_KE_TOAN)
    duoi = uuid4().hex[:6]
    thuoc = {}
    for ten, gia in (("A", 55_000), ("B", 105_000)):
        thuoc[ten] = await pool.fetchval(
            "INSERT INTO drug_catalog(clinic_id,name_base,name_raw,ma_hang,unit_price)"
            " VALUES($1::uuid,$2,$2,$3,$4) RETURNING id::text",
            manager.clinic_id,
            f"Thuốc {ten} {duoi}",
            f"T{ten}-{duoi}",
            gia,
        )
    cycle, ids = await _phieu(
        pool,
        q,
        manager,
        [
            ("prescription", f"Thuốc A {duoi}", "Hộp", 2, 55_000, thuoc["A"]),
            ("prescription", f"Thuốc B {duoi}", "Hộp", 1, 105_000, thuoc["B"]),
        ],
        kind="thuoc",
    )
    svc = ThueSuatService(pool)
    # Nhóm Thuốc 5% (nhập "5%"); riêng thuốc A đặt theo MẶT HÀNG 10%.
    await svc.dat_nhom(identity=ke_toan, nhom="THUOC", thue_suat="5%")
    await svc.dat_mat_hang(
        identity=ke_toan, loai="thuoc", id_=thuoc["A"], thue_suat="10"
    )
    ngay = now_vn().date().isoformat()
    xuat = XuatHoaDonService(pool)
    try:
        data = await xuat.xem(identity=ke_toan, tu=ngay, den=ngay)
        mine = [h for h in data["hoa_don"] if h["ma_phieu"] == ma_phieu(cycle)]
        theo_k = {h["thue_suat_ma"]: h for h in mine}
        assert set(theo_k) == {"10", "5"}
        # A: 110.000 / 1,10 = 100.000 + 10.000 · B: 105.000 / 1,05 = 100.000 + 5.000
        assert (theo_k["10"]["tien_hang"], theo_k["10"]["tien_thue"]) == (
            100_000,
            10_000,
        )
        assert (theo_k["5"]["tien_hang"], theo_k["5"]["tien_thue"]) == (100_000, 5_000)
        assert theo_k["10"]["dong"][0]["thue_nguon"] == "mat_hang"
        assert theo_k["5"]["dong"][0]["thue_nguon"] == "nhom"
        # Tổng khách trả giữ nguyên = tiền hàng + tiền thuế.
        assert sum(h["tong"] for h in mine) == 215_000
        assert sum(h["tien_hang"] + h["tien_thue"] for h in mine) == 215_000
        # stt duy nhất toàn file (mỗi thuế suất một hoá đơn riêng).
        stt = [h["stt"] for h in data["hoa_don"]]
        assert stt == list(range(1, len(stt) + 1))
        # Excel: K = SỐ, L = tiền thuế, thành tiền CHƯA thuế.
        ws: Any = load_workbook(io.BytesIO(xlsx_misa(mine))).active
        hang = {
            ws.cell(r, 11).value: r
            for r in range(DONG_DAU, ws.max_row + 1)
            if ws.cell(r, 11).value
        }
        r10 = hang[10]
        assert (ws.cell(r10, 12).value, ws.cell(r10, 17).value) == (10_000, 100_000)

        # Một hoá đơn: nhiều thuế suất → K trống, L = tổng thuế (khung).
        mot = await xuat.xem(identity=ke_toan, tu=ngay, den=ngay, che_do="mot")
        (m,) = [h for h in mot["hoa_don"] if h["ma_phieu"] == ma_phieu(cycle)]
        assert m["thue_suat_ma"] == "" and m["tien_thue"] == 15_000
        assert m["tong"] == 215_000 and len(m["dong"]) == 2

        # Sửa tay một dòng: thuốc A → KCT (thắng thuế mặt hàng 10).
        sua = await xuat.xem(
            identity=ke_toan,
            tu=ngay,
            den=ngay,
            ghi_de={ids[f"Thuốc A {duoi}"]: "KCT"},
        )
        kct = [
            h
            for h in sua["hoa_don"]
            if h["ma_phieu"] == ma_phieu(cycle) and h["thue_suat_ma"] == "KCT"
        ]
        assert len(kct) == 1 and kct[0]["tien_thue"] == 0
        assert kct[0]["dong"][0]["thanh_tien"] == 110_000

        # K2: bật công tắc; phiếu này không có giảm giá → vẫn một dòng mỗi thuốc.
        k2 = await xuat.xem(identity=ke_toan, tu=ngay, den=ngay, ghi_giam_gia=True)
        assert k2["ghi_giam_gia"] is True
    finally:
        await pool.execute(
            "DELETE FROM thue_suat_nhom WHERE clinic_id=$1::uuid AND nhom='THUOC'",
            manager.clinic_id,
        )


async def test_chua_dat_thue_xuat_y_nhu_cu_va_dem_dong_chua_thue(
    pool: asyncpg.Pool,
) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    duoi = uuid4().hex[:6]
    cycle, _ = await _phieu(
        pool, q, manager, [("exam", f"Khám {duoi}", None, 1, 300_000, None)]
    )
    ngay = now_vn().date().isoformat()
    data = await XuatHoaDonService(pool).xem(identity=manager, tu=ngay, den=ngay)
    (hd,) = [h for h in data["hoa_don"] if h["ma_phieu"] == ma_phieu(cycle)]
    assert hd["thue_suat_ma"] == "" and hd["tien_thue"] is None
    assert hd["dong"][0]["thanh_tien"] == 300_000 and hd["tong"] == 300_000
    assert data["dong_chua_thue"] >= 1
    ws: Any = load_workbook(io.BytesIO(xlsx_misa([hd]))).active
    assert ws.cell(DONG_DAU, 11).value is None and ws.cell(DONG_DAU, 12).value is None
    assert ws.cell(DONG_DAU, 17).value == 300_000


async def test_cong_quyen_dat_thue_va_dau_vao_rac(pool: asyncpg.Pool) -> None:
    svc = ThueSuatService(pool)
    khac = await nguoi_theo_lego(pool, "CSKH", ["bao_cao", "ds_benh_nhan"])
    with pytest.raises(SafetyGateError):
        await svc.xem(identity=khac)
    with pytest.raises(SafetyGateError):
        await svc.dat_nhom(identity=khac, nhom="DICH_VU", thue_suat="8")
    # Chỉ XEM tài chính (không có quyền đặt thuế): xem được, đặt thì không.
    ke_toan = await nguoi_theo_lego(pool, "CSKH", GOI_KE_TOAN)
    await pool.execute(
        "UPDATE capability_grant SET revoked_at=now(), revoked_by=staff_id"
        " WHERE staff_id=$1::uuid"
        " AND capability='accounting.tax.manage'",
        ke_toan.staff_id,
    )
    from clinicai.permissions import cache

    cache.quen(ke_toan.clinic_id, ke_toan.staff_id)
    assert (await svc.xem(identity=ke_toan))["duoc_dat"] is False
    with pytest.raises(SafetyGateError):
        await svc.dat_nhom(identity=ke_toan, nhom="DICH_VU", thue_suat="8")
    # Người đặt được: rác → 422 tiếng Việt, nhóm lạ → lỗi, trống = bỏ đặt.
    nguoi = await nguoi_theo_lego(pool, "CSKH", ["ke_toan"])
    for rac in ("abc", "-3", "101", "8%%"):
        with pytest.raises(ValidationError):
            await svc.dat_nhom(identity=nguoi, nhom="THUOC", thue_suat=rac)
    with pytest.raises(ValidationError):
        await svc.dat_nhom(identity=nguoi, nhom="KHAC", thue_suat="8")
    with pytest.raises(ValidationError):
        await svc.dat_mat_hang(
            identity=nguoi, loai="thuoc", id_=str(uuid4()), thue_suat="8"
        )
    # Ghi / xoá ở tầng MẶT HÀNG (một thuốc riêng của bài) — không đụng bảng nhóm
    # dùng chung với các bài xuất hoá đơn chạy song song.
    duoi = uuid4().hex[:6]
    thuoc = await pool.fetchval(
        "INSERT INTO drug_catalog(clinic_id,name_base,name_raw,unit_price)"
        " VALUES($1::uuid,$2,$2,1000) RETURNING id::text",
        nguoi.clinic_id,
        f"Thuốc đặt thuế {duoi}",
    )
    await svc.dat_mat_hang(identity=nguoi, loai="thuoc", id_=thuoc, thue_suat=" 8,5 ")
    (hang,) = (await svc.xem(identity=nguoi, tim=duoi))["mat_hang"]
    assert (hang["loai"], hang["thue_suat_ma"]) == ("thuoc", "8.5")
    await svc.dat_mat_hang(identity=nguoi, loai="thuoc", id_=thuoc, thue_suat="")
    (hang,) = (await svc.xem(identity=nguoi, tim=duoi))["mat_hang"]
    assert hang["thue_suat_ma"] is None


async def test_rang_buoc_db_thue_suat_hop_le(pool: asyncpg.Pool) -> None:
    for ma, ok in (
        ("8", True),
        ("0", True),
        ("100", True),
        ("8.25", True),
        ("KCT", True),
        ("KKKNT", True),
        (None, True),
        ("100.5", False),
        ("101", False),
        ("-1", False),
        ("8.255", False),
        ("kct", False),
        ("", False),
        ("abc", False),
    ):
        assert await pool.fetchval("SELECT thue_suat_hop_le($1)", ma) is ok, ma
    with pytest.raises(asyncpg.CheckViolationError):
        await pool.execute(
            "UPDATE service_price SET thue_suat_ma='abc'"
            " WHERE id=(SELECT id FROM service_price LIMIT 1)"
        )
