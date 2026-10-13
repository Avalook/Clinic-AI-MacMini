"""Excel downloads must remain editable, typed and safe to open."""

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from clinicai.services.excel_export import excel_table


def test_excel_table_preserves_types_and_literal_formulas():
    content = excel_table(
        [
            (
                "Chi tiết",
                ["Mã", "Tiền", "Ngày", "Lúc"],
                [
                    [
                        '=HYPERLINK("https://example.com")',
                        Decimal("250000"),
                        date(2026, 10, 12),
                        datetime(2026, 10, 12, 12, 30, tzinfo=timezone.utc),
                    ],
                    ["00123", -50000, None, None],
                ],
            )
        ]
    )
    ws = load_workbook(BytesIO(content)).active
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value.startswith("=HYPERLINK")
    assert ws["A3"].value == "00123"
    assert ws["B2"].data_type == "n"
    assert ws["B3"].value == -50000
    assert ws["C2"].data_type == "d"
    assert ws["D2"].value.hour == 19
    assert ws.freeze_panes == "A2"
    assert not ws.protection.sheet
    assert not ws.merged_cells
    assert next(iter(ws.tables.values())).ref == "A1:D3"


def test_excel_table_empty_data_still_has_filterable_headers():
    ws = load_workbook(BytesIO(excel_table([("Dữ liệu", ["Mã", "Tiền"], [])]))).active
    assert ws.max_row == 1
    assert next(iter(ws.tables.values())).ref == "A1:B1"


def test_xlsx_download_routes_keep_authenticated_dependencies():
    from clinicai.api.identity import get_current_identity
    from clinicai.api.v1.routers import cashier, pharmacy, reports

    for router, path in (
        (reports.router, "/reports/cuoi-ngay.xlsx"),
        (cashier.router, "/cashier/lich-su.xlsx"),
        (pharmacy.router, "/pharmacy/xuat-nhap-ton.xlsx"),
    ):
        route = next(r for r in router.routes if getattr(r, "path", "") == path)
        calls = [dep.call for dep in route.dependant.dependencies]
        assert get_current_identity in calls or pharmacy._DOC in calls


def _report():
    from clinicai.services.bao_cao_cuoi_ngay_service import gom_bao_cao

    return gom_bao_cao(
        tu=date(2026, 10, 12),
        den=date(2026, 10, 13),
        lan_thu=[
            {
                "id": "1",
                "amount": 300000,
                "kind": "dich_vu",
                "status": "PAID",
                "method": "CASH",
                "paid_at": datetime.fromisoformat("2026-10-12T19:00:00+07:00"),
            }
        ],
        hoan=[
            {
                "amount": 50000,
                "status": "COMPLETED",
                "kind": "dich_vu",
                "method": "CASH",
                "created_at": datetime.fromisoformat("2026-10-13T19:00:00+07:00"),
            }
        ],
        dong=[],
        doi_tac=[],
        so_luot_kham=1,
    )


def test_report_summary_matches_filtered_net_not_original_receipts():
    from clinicai.services.excel_report_export import xlsx_tom_tat

    report = _report()
    report["khach_con_no"] = "hidden"
    report["tai_chinh"] = {"doanh_thu": 990000}
    ws = load_workbook(BytesIO(xlsx_tom_tat(report))).active
    rows = dict(list(ws.values)[1:])
    assert rows["Tổng thu"] == 250000
    assert rows["Dịch vụ"] == 250000
    assert "Doanh thu" not in rows
    assert "hidden" not in str(rows)
    detailed = load_workbook(BytesIO(xlsx_tom_tat(report, day_du=True))).active
    assert dict(list(detailed.values)[1:])["Doanh thu"] == 990000


def test_report_detail_and_inventory_keep_tables_and_numeric_data():
    from clinicai.services.excel_export import (
        excel_date,
        excel_datetime,
        xlsx_inventory,
    )
    from clinicai.services.excel_report_export import xlsx_bao_cao, xlsx_hang_hoa

    report = _report()
    report["thuoc_theo_khach"] = {"khach": [], "tieu_hao": []}
    report["theo_co_so"] = [
        {"ten": "A", "khach": report["khach"], "tong": report["tong"]}
    ]
    report["tai_chinh"] = {"doanh_thu": 300000}
    book = load_workbook(BytesIO(xlsx_bao_cao(report)))
    assert "Theo cơ sở" in book.sheetnames
    assert "Thuốc theo khách" in book.sheetnames
    assert all(ws.tables for ws in book.worksheets)
    assert book["Theo ngày"]["A2"].data_type == "d"
    assert book["Hoàn hủy"]["B2"].value.hour == 19
    assert load_workbook(BytesIO(xlsx_hang_hoa(report, "rác"))).active.tables
    inventory = {
        "dong": [
            {
                "ma_hang": "001",
                "ten": "=1+1",
                "ton_dau": Decimal("2.5"),
                "gia_tri_ton": None,
            }
        ]
    }
    ws = load_workbook(BytesIO(xlsx_inventory(inventory))).active
    assert ws["A2"].value == "001"
    assert ws["B2"].data_type == "s"
    assert ws["D2"].value == 2.5
    assert excel_date("rác") is None
    assert excel_datetime("rác") is None
    assert excel_datetime(datetime(2026, 10, 12)) == datetime(2026, 10, 12)


def test_cashier_export_keeps_refunds_as_numbers_and_timestamps():
    from clinicai.services.quay_thu_service import xlsx_lich_su

    group = {
        "ten": "=SUM(1)",
        "ma_bn": "001",
        "su_kien": [
            {
                "loai": "hoan",
                "so_tien": -50000,
                "luc": "2026-10-13T12:00:00Z",
                "cho": True,
            }
        ],
    }
    ws = load_workbook(BytesIO(xlsx_lich_su([group]))).active
    assert ws["D2"].data_type == "s"
    assert ws["K2"].value == -50000
    assert ws["A2"].value.hour == 19


def test_cashier_discount_export_preserves_net_and_does_not_repeat_discount():
    from clinicai.services.quay_thu_service import gom_theo_khach, xlsx_lich_su

    cycle = {
        "id": "12345678-1234-1234-1234-123456789abc",
        "visit_id": "visit",
        "amount": 800000,
        "original_amount": 1000000,
        "discount_amount": 200000,
        "discount_reason": "=khách thân thiết",
        "discount_by_name": "Nhân viên A",
        "discount_type": "PERCENT",
        "discount_value": Decimal("20"),
        "status": "VOIDED",
    }
    groups = gom_theo_khach([cycle], [], {}, {})
    receipt = groups[0]["phieu"][0]
    assert receipt["discount_amount"] == 200000
    assert receipt["tong"] == 800000
    ws = load_workbook(BytesIO(xlsx_lich_su(groups))).active
    headers = {cell.value: cell.column for cell in ws[1]}
    rows = [dict(zip(headers, row)) for row in list(ws.values)[1:]]
    paid = next(row for row in rows if row["Loại"] == "Thu")
    canceled = next(row for row in rows if row["Loại"] != "Thu")
    assert paid["Số tiền"] == 800000
    assert paid["Trước giảm giá"] == 1000000
    assert paid["Giảm giá"] == 200000
    assert paid["Người giảm giá"] == "Nhân viên A"
    assert paid["Lý do giảm giá"] == "=khách thân thiết"
    assert canceled["Giảm giá"] is None
    assert canceled["Số tiền"] == -800000


def test_cashier_legacy_discount_metadata_defaults_without_changing_totals():
    from clinicai.services.quay_thu_service import gom_theo_khach, xlsx_lich_su

    cycle = {
        "id": "12345678-1234-1234-1234-123456789abc",
        "visit_id": "visit",
        "amount": 800000,
        "status": "PAID",
    }
    group = gom_theo_khach([cycle], [], {}, {})
    ws = load_workbook(BytesIO(xlsx_lich_su(group))).active
    row = dict(zip(list(ws.values)[0], list(ws.values)[1]))
    assert row["Trước giảm giá"] == 800000
    assert row["Giảm giá"] == 0
    assert group[0]["con_lai"] == 800000


def test_report_discount_metadata_keeps_actual_summary_and_audit_details():
    from clinicai.services.bao_cao_cuoi_ngay_service import gom_bao_cao
    from clinicai.services.excel_report_export import xlsx_bao_cao, xlsx_tom_tat

    cycle = {
        "id": "discounted",
        "amount": 800000,
        "kind": "dich_vu",
        "status": "PAID",
        "method": "CASH",
        "paid_at": datetime(2026, 10, 12),
        "original_amount": 1000000,
        "discount_amount": 200000,
        "discount_reason": "Khách thân thiết",
        "discount_by_name": "Nhân viên A",
        "discount_type": "PERCENT",
        "discount_value": 20,
    }
    report = gom_bao_cao(
        tu=date(2026, 10, 12),
        den=date(2026, 10, 12),
        lan_thu=[cycle],
        hoan=[],
        dong=[],
        doi_tac=[],
        so_luot_kham=1,
    )
    assert report["tong"]["thuc_thu"] == 800000
    ws = load_workbook(BytesIO(xlsx_bao_cao(report)))["Giảm giá"]
    row = dict(zip(list(ws.values)[0], list(ws.values)[1]))
    assert row["Giảm giá"] == 200000
    assert row["Sau giảm giá"] == 800000
    assert row["Người giảm giá"] == "Nhân viên A"
    simple = load_workbook(BytesIO(xlsx_tom_tat(report)))
    assert dict(list(simple.active.values)[1:])["Tổng thu"] == 800000
    assert "Giảm giá" not in simple.sheetnames


@pytest.mark.asyncio
async def test_report_xlsx_endpoint_uses_authorized_scope(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from clinicai.api.v1.routers import reports
    from clinicai.core.exceptions import SafetyGateError

    authorize = AsyncMock(
        return_value=({"tai_chinh": True, "legacy": True}, "allowed-location")
    )
    report = AsyncMock(return_value=_report())
    monkeypatch.setattr(
        reports, "ReportAccessService", lambda _: SimpleNamespace(authorize=authorize)
    )
    monkeypatch.setattr(
        reports, "BaoCaoCuoiNgayService", lambda _: SimpleNamespace(bao_cao=report)
    )
    params = dict(
        tu="2026-10-12",
        den="2026-10-13",
        loai=None,
        ca=None,
        co_so="requested",
        mau="tom_tat",
        identity=SimpleNamespace(),
        pool=None,
    )
    response = await reports.bao_cao_cuoi_ngay_xlsx(**params)
    assert response.body[:2] == b"PK"
    assert response.headers["content-disposition"].endswith('.xlsx"')
    assert response.headers["cache-control"] == "private, no-store"
    assert report.call_args.kwargs["co_so"] == "allowed-location"
    data = _report()
    data["tai_chinh"] = {"doanh_thu": 990000}
    report.return_value = data
    params["mau"] = "tom_tat_day_du"
    response = await reports.bao_cao_cuoi_ngay_xlsx(**params)
    assert "Doanh thu" not in dict(
        list(load_workbook(BytesIO(response.body)).active.values)[1:]
    )
    authorize.return_value = (
        {"tai_chinh": True, "legacy": True, "tai_chinh_day_du": True},
        "allowed-location",
    )
    response = await reports.bao_cao_cuoi_ngay_xlsx(**params)
    assert (
        dict(list(load_workbook(BytesIO(response.body)).active.values)[1:])["Doanh thu"]
        == 990000
    )
    authorize.return_value = ({"tai_chinh": False, "legacy": False}, "allowed-location")
    with pytest.raises(SafetyGateError):
        await reports.bao_cao_cuoi_ngay_xlsx(**params)
