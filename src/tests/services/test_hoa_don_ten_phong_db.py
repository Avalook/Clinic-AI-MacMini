"""Danh sách hoá đơn hiện TÊN PHÒNG dịch vụ — đúng nguồn phiếu thu in (30/10/2026).

    scripts/test-nhanh.sh src/tests/services/test_hoa_don_ten_phong_db.py

Dòng dịch vụ đã xếp phòng: màn /hoa-don và file Excel ghi cùng chữ phòng như
phiếu thu 80mm ("Phòng X · BS Y"). Dòng bỏ sau thu không ghi phòng (như bản in
lại). Bản xuất MISA không đổi.
"""

from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from typing import Any

import asyncpg
import pytest
from openpyxl import load_workbook

from clinicai.api.identity import ClinicRole
from clinicai.core.clock import now_vn
from clinicai.services.danh_sach_hoa_don_service import (
    DanhSachHoaDonService,
    xlsx_hoa_don,
)
from clinicai.services.hoa_don_projection import nhan_phong
from clinicai.services.hoan_tac_service import HoanTacService
from clinicai.services.quay_thu_service import QuayThuService
from tests.chay_nguoi_dua_tin import chay_hanh_trinh
from tests.services.test_check_in_lai_sau_hoan_tac_db import (  # noqa: F401
    pool,
)
from tests.services.test_thu_tien_xep_phong_mang_sang_db import (
    _benh_nhan,
    _check_in,
    _chon,
    _dung,
    _kham_va_chi_dinh,
    _thu,
)

pytestmark = [pytest.mark.db, pytest.mark.asyncio]


def _cot(ws: Any, ten: str) -> list[Any]:
    dau = [c.value for c in ws[1]]
    i = dau.index(ten)
    return [r[i].value for r in ws.iter_rows(min_row=2)]


async def test_hoa_don_ghi_phong_nhu_phieu_thu_va_excel(
    pool: asyncpg.Pool,  # noqa: F811
) -> None:
    ca = await _dung(pool)
    visit = await _check_in(pool, ca, await _benh_nhan(pool, ca), ca.loai_kham)
    _con, order = await _kham_va_chi_dinh(pool, ca, visit)
    await _chon(pool, ca, visit, [order])
    await _thu(pool, visit, ca.thu_ngan)
    await chay_hanh_trinh(pool)
    ten_phong = await pool.fetchval(
        "SELECT r.name FROM service_order o JOIN clinic_room r ON r.id = o.room_id"
        " WHERE o.id = $1::uuid",
        order,
    )
    assert ten_phong

    quan_ly = replace(ca.thu_ngan, role=ClinicRole.MANAGEMENT)
    day = now_vn().date().isoformat()
    svc = DanhSachHoaDonService(pool)

    async def phieu_cua_luot() -> dict[str, Any]:
        data = await svc.xem(identity=quan_ly, tu=day, den=day)
        [p] = [p for k in data["khach"] for p in k["phieu"] if p["visit_id"] == visit]
        return dict(p)

    p = await phieu_cua_luot()
    [dv] = [d for d in p["dong"] if d["source_type"] == "service_order"]
    # Cùng chữ phiếu thu in cho đúng phiếu ấy.
    in_ra = await QuayThuService(pool).phieu(
        identity=ca.thu_ngan, id_=p["id"], loai="thu"
    )
    [dong_in] = [d for d in in_ra["dong"] if d.get("order_id") == order]
    assert dv["phong"] == nhan_phong(dong_in["phong"])
    assert dv["phong"].startswith(ten_phong)
    # Dòng không phải dịch vụ (tiền khám) không có phòng.
    assert all(
        d["phong"] is None for d in p["dong"] if d["source_type"] != "service_order"
    )

    data = await svc.xem(identity=quan_ly, tu=day, den=day)
    book = load_workbook(BytesIO(xlsx_hoa_don(data)))
    ws = book["Dòng phiếu"]
    hang = dict(zip(_cot(ws, "Dịch vụ / thuốc"), _cot(ws, "Phòng"), strict=True))
    assert hang[dv["ten"]] == dv["phong"]

    # Bỏ chỉ định sau thu → dòng "đã bỏ", không ghi phòng (như bản in lại).
    await HoanTacService(pool).huy_chi_dinh(
        order_id=order, identity=ca.bac_si, xac_nhan=True, ly_do="Khách đổi ý"
    )
    [dv] = [
        d
        for d in (await phieu_cua_luot())["dong"]
        if d["source_type"] == "service_order"
    ]
    assert dv["da_bo"] is True and dv["phong"] is None
