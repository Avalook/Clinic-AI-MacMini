"""Theo dõi thuốc (31/10/2026): số "đang ở đâu" khớp tồn lô, dòng thời gian đủ mọi
loại việc, cơ sở khác không lọt, đầu vào rác trả rỗng không ném.

Chạy lệnh thật từ đầu tới cuối một vòng đời thuốc ở cơ sở riêng (fixture `co`).
"""

from __future__ import annotations

# ruff: noqa: F811 — fixture q / co được import từ CP1 / G7.
import dataclasses
import io
import uuid
from decimal import Decimal
from typing import Any

import httpx
import openpyxl
import pytest

from clinicai.api.identity import _resolve_identity
from clinicai.api.v1.routers.pharmacy import router as pharmacy_router
from clinicai.core.clock import hom_nay_vn
from clinicai.core.database import get_db_pool
from clinicai.main import app
from clinicai.services import kho_thuoc_service as kho
from clinicai.services import kiem_ke_ca_service as kk
from clinicai.services import lo_thuoc_service
from clinicai.services import theo_doi_thuoc_service as td
from clinicai.services.excel_export import xlsx_theo_doi_thuoc
from clinicai.services.pharmacy_service import PharmacyService
from tests.services.test_kho_kiotviet_db import _thuoc_moi
from tests.services.test_kho_quay_db import (  # noqa: F401
    Co,
    _bat,
    _chuyen,
    _id_dich,
    _nhap_kho,
    co,
)
from tests.services.test_tien_thuoc_cp1_db import Quay, _nhap_lo, q  # noqa: F401
from tests.services.test_tien_thuoc_cp3_db import _dong_da_xac_dinh, _thu
from tests.services.test_tien_thuoc_cp6_dinh_chinh_db import _lan

pytest_plugins = ["tests.services.test_luot_kham_service_db"]
pytestmark = [pytest.mark.db, pytest.mark.asyncio]
D = Decimal


@pytest.fixture(autouse=True)
def _thu_khong_cho_kho(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("CLINICAI_DRUG_PAYMENT_REQUIRES_INVENTORY", "0")


@dataclasses.dataclass
class Vong:
    co: Co
    drug: str
    ma_lo: str
    kho_lo: str
    quay: str


async def _vong_doi(co: Co) -> Vong:
    """Một thuốc đi đủ: nhập, kiểm, chuyển, giao, trả, huỷ, gán lô, hoàn tác."""
    pool = co.q.pool
    svc = PharmacyService(pool)
    rx, drug = await _dong_da_xac_dinh(co.q, 8)
    await _bat(co)
    ma_lo = f"LO-{uuid.uuid4().hex[:10]}"
    kho_lo = await _nhap_kho(co, drug, 30, ma=ma_lo)
    await kho.kiem_kho(
        pool,
        identity=co.ds,
        khoa_gui=uuid.uuid4().hex,
        dong=[{"drug_batch_id": kho_lo, "thuc_te": 29, "ly_do": "Vỡ lúc xếp kệ"}],
    )
    quay = _id_dich(await _chuyen(co, kho_lo, 12))

    # Kiểm ca lệch −1 rồi hoàn tác.
    phieu = await kk.mo_phieu(
        pool, identity=co.ds, ngay=hom_nay_vn().isoformat(), ca="TOI", loai="CUOI"
    )
    await kk.luu_dem(
        pool,
        identity=co.ds,
        phieu_id=phieu["id"],
        dong=[{"drug_catalog_id": drug, "don_vi": "viên", "dem": 11}],
    )
    await kk.chot(pool, identity=co.ds, phieu_id=phieu["id"])
    await kk.dieu_chinh_theo_kiem_ke(
        pool, identity=co.ds, phieu_id=phieu["id"], ly_do="Đếm lại cuối ca"
    )
    await kk.hoan_tac_dieu_chinh(pool, identity=co.ds, phieu_id=phieu["id"])

    await svc.dieu_chinh(
        identity=co.ds, drug_batch_id=quay, so_luong=1, ly_do="Tìm thấy một vỉ"
    )
    await svc.huy(identity=co.ds, drug_batch_id=kho_lo, so_luong=1, ly_do="Vỡ")

    # Huỷ lô hết hạn rồi hoàn tác: không tính vào "đã huỷ".
    lo2 = await _nhap_kho(co, drug, 5)
    await pool.execute(
        "UPDATE drug_batch SET expiry_date = current_date - 1 WHERE id = $1::uuid", lo2
    )
    huy = await lo_thuoc_service.huy_lo(
        pool,
        identity=co.ds,
        drug_batch_id=lo2,
        ly_do="Hết hạn",
        khoa_gui=str(uuid.uuid4()),
    )
    await lo_thuoc_service.hoan_huy_lo(
        pool, identity=co.ds, inventory_txn_id=huy["inventory_txn_id"]
    )

    # Giao 3, trả 3 lần: một bán lại, một huỷ, một còn chờ kiểm.
    await _thu(co.q)
    await svc.cap_phat(
        identity=co.ds, prescription_id=rx, drug_batch_id=quay, so_luong=3
    )
    txn = await pool.fetchval(
        "SELECT id::text FROM inventory_txn WHERE ref_id = $1::uuid"
        " AND txn_type = 'DISPENSE'",
        rx,
    )
    tra = [
        (
            await svc.khach_tra_thuoc(
                identity=co.ds,
                dispense_txn_id=txn,
                so_luong=1,
                ly_do="Khách đổi ý",
                ket_cuc_tien="GIU_LAI",
                ly_do_tien="Giữ lại tiền thuốc",
            )
        )["drug_return_id"]
        for _ in range(3)
    ]
    await svc.xu_ly_thuoc_tra(
        identity=co.ds, drug_return_id=tra[0], quyet_dinh="BAN_LAI", ly_do="Còn nguyên"
    )
    await svc.xu_ly_thuoc_tra(
        identity=co.ds, drug_return_id=tra[1], quyet_dinh="HUY", ly_do="Đã mở hộp"
    )

    # Rút hết phần bán được về kho → giao nữa thì chờ gán lô, kho chưa trừ.
    con = await pool.fetchval(
        "SELECT public.drug_batch_kha_dung($1::uuid, $2::uuid)", co.ds.clinic_id, quay
    )
    await _chuyen(co, quay, con, huong="QUAY_KHO")
    for _ in range(2):
        await svc.cap_phat(
            identity=co.ds, prescription_id=rx, drug_batch_id=None, so_luong=1
        )
    dong = await pool.fetchval(
        "SELECT id::text FROM thuoc_giao_chua_gan_lo WHERE prescription_id = $1::uuid"
        " ORDER BY giao_luc LIMIT 1",
        rx,
    )
    await svc.gan_lo_da_giao(identity=co.ds, dong_id=dong, drug_batch_id=kho_lo)

    k = await _chuyen(co, kho_lo, 2)
    await kho.hoan_tac_chuyen(
        pool, identity=co.ds, khoa_gui=uuid.uuid4().hex, phieu_id=k["id"]
    )
    return Vong(co, drug, ma_lo, kho_lo, quay)


async def _xem(co: Co, drug: str, **kw: Any) -> dict[str, Any]:
    return await td.theo_doi_thuoc(co.q.pool, identity=co.ds, thuoc=drug, **kw)


async def _tat_ca(co: Co, drug: str, **kw: Any) -> list[dict[str, Any]]:
    kq = await td.xuat_theo_doi(co.q.pool, identity=co.ds, thuoc=drug, **kw)
    return list(kq["dong"])


async def test_vong_doi_day_du(co: Co) -> None:
    v = await _vong_doi(co)
    kq = await _xem(co, v.drug)
    dong = await _tat_ca(co, v.drug)

    # ── Đủ mọi loại việc, chữ thường, mới nhất trước ──
    assert {d["ma"] for d in dong} == {
        "RECEIVE",
        "KIEM_KE_LECH",
        "CHUYEN_RA_QUAY",
        "TRA_VE_KHO",
        "HOAN_TAC_CHUYEN",
        "DISPENSE",
        "GAN_LO",
        "RETURN_RECEIVED",
        "NHAP_LAI_BAN",
        "DISCARD",
        "HOAN_TAC_HUY",
        "HOAN_TAC_KIEM_KE",
        "ADJUST",
    }
    assert {d["hanh_dong"] for d in dong} <= set(td.HANH_DONG.values())
    assert [d["luc"] for d in dong] == sorted((d["luc"] for d in dong), reverse=True)
    assert kq["tong"] == len(dong)
    assert all(d["nguoi_lam"] for d in dong)

    # ── Tồn trước → sau: cộng đúng, dòng mới nhất mỗi hàng = tồn hiện tại ──
    hang = await co.q.pool.fetch(
        "SELECT batch_code, noi_chua, quantity_on_hand FROM drug_batch"
        " WHERE drug_catalog_id = $1::uuid AND location_id = $2::uuid",
        v.drug,
        co.loc,
    )
    vet: dict[tuple[str, str], list[tuple[Decimal, Decimal]]] = {}
    for d in dong:
        if d["ton_sau"] is None:
            continue
        doi = -d["so_luong"] if d["la_chuyen"] else d["so_luong"]
        assert d["ton_truoc"] + doi == d["ton_sau"]
        vet.setdefault((d["so_lo"], d["noi"]), []).append(
            (d["ton_truoc"], d["ton_sau"])
        )
        if d["la_chuyen"]:
            assert d["ton_den_truoc"] + d["so_luong"] == d["ton_den_sau"]
            vet.setdefault((d["so_lo"], d["noi_den"]), []).append(
                (d["ton_den_truoc"], d["ton_den_sau"])
            )
    for h in hang:
        chuoi = vet[(h["batch_code"], h["noi_chua"])]
        assert chuoi[0][1] == h["quantity_on_hand"]
        assert chuoi[-1][0] == 0
        # Liền mạch: tồn trước của dòng mới = tồn sau của dòng ngay trước nó.
        assert all(a[0] == b[1] for a, b in zip(chuoi, chuoi[1:], strict=False))

    # ── Đang ở đâu khớp tồn lô ──
    ton = {
        n: sum((h["quantity_on_hand"] for h in hang if h["noi_chua"] == n), D(0))
        for n in ("KHO", "QUAY")
    }
    [t] = kq["dang_o_dau"]["tong"]
    assert (t["don_vi"], t["o_kho"], t["o_quay"]) == ("viên", ton["KHO"], ton["QUAY"])
    assert t["tong"] == ton["KHO"] + ton["QUAY"]
    assert (t["cho_gan_lo"], t["cho_kiem"], t["da_huy"]) == (1, 1, 2)
    # Chờ kiểm đã nằm trong tồn; chờ gán chưa trừ nên trên kệ = tồn − chờ gán.
    assert t["tren_ke"] == t["tong"] - t["cho_gan_lo"]
    lo = kq["dang_o_dau"]["lo"]
    assert sum(x["o_kho"] for x in lo) == ton["KHO"]
    assert sum(x["o_quay"] for x in lo) == ton["QUAY"]
    chinh = next(x for x in lo if x["so_lo"] == v.ma_lo)
    assert chinh["han_dung"] is not None and chinh["het_han"] is False
    assert any(x["het_han"] for x in lo)  # lô đã cho hết hạn ở trên

    # ── Từ đâu → đến đâu, khách, phiếu, lý do ──
    def mot(ma: str) -> dict[str, Any]:
        return next(d for d in dong if d["ma"] == ma)

    assert (mot("RECEIVE")["tu"], mot("RECEIVE")["den"]) == ("Nhà cung cấp", "Kho")
    assert (mot("CHUYEN_RA_QUAY")["tu"], mot("CHUYEN_RA_QUAY")["den"]) == (
        "Kho",
        "Quầy",
    )
    assert mot("CHUYEN_RA_QUAY")["ma_phieu"].startswith("CK-")
    assert mot("KIEM_KE_LECH")["den"] is None
    assert any(
        (d["ma_phieu"] or "").startswith("Kiểm cuối ca tối ")
        for d in dong
        if d["ma"] == "KIEM_KE_LECH"
    )
    assert mot("HOAN_TAC_KIEM_KE")["ma_phieu"].startswith("Kiểm cuối ca tối ")
    tra = mot("RETURN_RECEIVED")
    assert (tra["tu"], tra["den"]) == ("Khách", "Chờ kiểm")
    assert tra["khach"] and tra["ma_bn"] and tra["visit_id"] == co.q.visit_id
    assert tra["ly_do"] == "Khách đổi ý"
    bl = mot("NHAP_LAI_BAN")
    assert bl["ton_sau"] is None and bl["so_luong"] == 1
    assert (bl["tu"], bl["den"]) == ("Chờ kiểm", "Quầy")
    assert bl["ly_do"] == "Còn nguyên" and bl["khach"] == tra["khach"]
    huy_tra = next(d for d in dong if d["ma"] == "DISCARD" and d["tu"] == "Chờ kiểm")
    assert huy_tra["den"] == "Đã huỷ" and "Đã mở hộp" in huy_tra["ly_do"]
    assert mot("HOAN_TAC_HUY")["tu"] == "Đã huỷ"

    # ── Chờ gán lô hiện dù chưa có dòng sổ; đã gán thì thành dòng sổ, giữ giờ giao ──
    [cho] = [d for d in dong if d["chua_tru_kho"]]
    assert cho["ma"] == "DISPENSE" and cho["hanh_dong"] == "Giao cho khách"
    assert cho["so_lo"] is None and cho["ton_sau"] is None and cho["so_luong"] == -1
    assert "Chưa trừ kho" in cho["ly_do"] and cho["khach"] and cho["nguoi_lam"]
    gan = mot("GAN_LO")
    assert gan["ton_sau"] is not None and gan["ly_do"].startswith("Giao khách lúc ")

    # ── Excel gồm đủ dòng, số vẫn là số ──
    bao_cao = await td.xuat_theo_doi(co.q.pool, identity=co.ds, thuoc=v.drug)
    so_sach = openpyxl.load_workbook(io.BytesIO(xlsx_theo_doi_thuoc(bao_cao)))
    assert so_sach["Dòng thời gian"].max_row == len(dong) + 1
    assert so_sach["Đang ở đâu"].max_row == len(lo) + 1


async def test_loc_va_phan_trang(co: Co) -> None:
    v = await _vong_doi(co)
    dong = await _tat_ca(co, v.drug)

    for k, gia, dung in (
        ("hanh_dong", "DISPENSE", lambda d: d["ma"] == "DISPENSE"),
        ("noi", "KHO", lambda d: "KHO" in (d["noi"], d["noi_den"])),
        ("lo", v.ma_lo, lambda d: d["so_lo"] == v.ma_lo),
    ):
        kq = await _xem(co, v.drug, **{k: gia})
        assert kq["tong"] == len([d for d in dong if dung(d)]) > 0
        assert all(dung(d) for d in kq["dong"])

    hom_nay = hom_nay_vn().isoformat()
    assert (await _xem(co, v.drug, tu=hom_nay, den=hom_nay))["tong"] == len(dong)
    assert (await _xem(co, v.drug, den="2021-01-01"))["tong"] == 0
    # Đảo hai đầu ngày vẫn ra đúng khoảng.
    assert (await _xem(co, v.drug, tu=hom_nay, den="2021-01-01"))["tong"] == len(dong)

    bo_loc = (await _xem(co, v.drug))["loc"]
    assert {h["ma"] for h in bo_loc["hanh_dong"]} == {d["ma"] for d in dong}
    [nguoi] = bo_loc["nguoi"]
    assert nguoi["id"] == co.ds.staff_id
    assert (await _xem(co, v.drug, nguoi=nguoi["id"]))["tong"] == len(dong)

    # Phân trang: không trùng, không sót, tổng không đổi theo trang.
    trang = [await _xem(co, v.drug, trang=n, so_dong=6) for n in range(1, 8)]
    gop = [d["id"] for t in trang for d in t["dong"]]
    assert gop == [d["id"] for d in dong]
    assert {t["tong"] for t in trang} == {len(dong)}
    assert (await _xem(co, v.drug, trang=9999))["dong"] == []


async def test_ton_khong_doi_khi_loc_hay_phan_trang(co: Co) -> None:
    v = await _vong_doi(co)
    goc = {d["id"]: d for d in await _tat_ca(co, v.drug)}

    def khop(ds: list[dict[str, Any]]) -> None:
        cot = ("ton_truoc", "ton_sau", "ton_den_truoc", "ton_den_sau")
        assert ds
        for d in ds:
            assert [d[c] for c in cot] == [goc[d["id"]][c] for c in cot]

    khop((await _xem(co, v.drug, hanh_dong="DISPENSE"))["dong"])
    khop((await _xem(co, v.drug, noi="QUAY", lo=v.ma_lo))["dong"])
    khop((await _xem(co, v.drug, trang=2, so_dong=3))["dong"])


async def test_chuyen_kho_gop_mot_su_kien_hoan_tac_hien_ro(co: Co) -> None:
    drug = await _thuoc_moi(co.q)
    await _bat(co)
    lo = await _nhap_kho(co, drug, 10)
    k = await _chuyen(co, lo, 4, ly_do="Quầy sắp hết")
    await kho.hoan_tac_chuyen(
        co.q.pool, identity=co.ds, khoa_gui=uuid.uuid4().hex, phieu_id=k["id"]
    )
    hoan, ra, nhap = await _tat_ca(co, drug)  # hai dòng sổ mỗi lần chuyển → một dòng
    assert nhap["ma"] == "RECEIVE"
    assert ra["ma"] == "CHUYEN_RA_QUAY" and ra["hanh_dong"] == "Chuyển ra quầy"
    assert (ra["tu"], ra["den"], ra["so_luong"], ra["la_chuyen"]) == (
        "Kho",
        "Quầy",
        4,
        True,
    )
    assert (ra["ton_truoc"], ra["ton_sau"]) == (10, 6)
    assert (ra["ton_den_truoc"], ra["ton_den_sau"]) == (0, 4)
    assert ra["ly_do"].startswith("Quầy sắp hết · Đã hoàn tác bằng phiếu CK-")
    assert (
        hoan["ma"] == "HOAN_TAC_CHUYEN" and hoan["hanh_dong"] == "Hoàn tác chuyển kho"
    )
    assert (hoan["tu"], hoan["den"], hoan["so_luong"]) == ("Quầy", "Kho", 4)
    assert (hoan["ton_truoc"], hoan["ton_sau"]) == (4, 0)
    assert (hoan["ton_den_truoc"], hoan["ton_den_sau"]) == (6, 10)
    assert (
        hoan["ly_do"].startswith("Hoàn tác phiếu CK-")
        and ra["ma_phieu"] in hoan["ly_do"]
    )
    # Lọc theo nơi: việc chuyển thuộc cả nơi đi lẫn nơi đến.
    assert (await _xem(co, drug, noi="QUAY"))["tong"] == 2
    assert (await _xem(co, drug, noi="KHO"))["tong"] == 3
    so_dong = await co.q.pool.fetchval(
        "SELECT count(*) FROM inventory_txn t"
        " JOIN drug_batch b ON b.id = t.drug_batch_id"
        " WHERE b.drug_catalog_id = $1::uuid",
        drug,
    )
    assert so_dong == 5  # sổ vẫn đủ: chỉ gộp lúc hiện


async def _lo_tay(co: Co, drug: str, don_vi: str, so: int, ly_do: str = "Nhập") -> str:
    """Lô đơn vị bất kỳ (lệnh nhập có quy đổi không cho tạo) + dòng sổ nhập."""
    lo = str(
        await co.q.pool.fetchval(
            "INSERT INTO drug_batch (clinic_id, location_id, drug_catalog_id,"
            " batch_code, expiry_date, quantity_on_hand, unit, noi_chua, received_at,"
            " ma_tem) VALUES ($1::uuid, $2::uuid, $3::uuid, $4, '2099-12-31', 0, $5,"
            " 'QUAY', now(), 'LT' || lpad(nextval('drug_batch_tem_seq')::text, 8, '0'))"
            " RETURNING id",
            co.ds.clinic_id,
            co.loc,
            drug,
            f"LO-{uuid.uuid4().hex[:10]}",
            don_vi,
        )
    )
    await _ghi_so(co, lo, "RECEIVE", so, ly_do)
    return lo


async def _ghi_so(
    co: Co, lo: str, loai: str, so: int, ly_do: str | None, ref: str | None = None
) -> None:
    await co.q.pool.execute(
        "INSERT INTO inventory_txn (clinic_id, drug_batch_id, txn_type, quantity,"
        " reason, ref_type, ref_id, performed_by_staff_id)"
        " VALUES ($1::uuid, $2::uuid, $3, $4, $5, $6, $7::uuid, $8::uuid)",
        co.ds.clinic_id,
        lo,
        loai,
        so,
        ly_do,
        "prescription" if ref else "manual",
        ref,
        co.ds.staff_id,
    )


async def test_nhieu_don_vi_lo_khong_cong_lan_don_vi_la_khong_vo(co: Co) -> None:
    drug = await _thuoc_moi(co.q)
    await _lo_tay(co, drug, "hộp", 2)
    await _lo_tay(co, drug, "Hộp ", 1)  # cùng đơn vị, khác cách viết
    await _lo_tay(co, drug, "gói", 30)
    await _lo_tay(co, drug, "túi 3 vỉ x 10", 4)

    kq = await _xem(co, drug)
    tong = {t["don_vi"].lower(): t for t in kq["dang_o_dau"]["tong"]}
    assert set(tong) == {"hộp", "gói", "túi 3 vỉ x 10"}
    assert [tong[k]["tong"] for k in ("hộp", "gói", "túi 3 vỉ x 10")] == [3, 30, 4]
    assert sum(t["tong"] for t in tong.values()) == sum(
        x["o_quay"] for x in kq["dang_o_dau"]["lo"]
    )
    assert {d["don_vi"].strip().lower() for d in kq["dong"]} == set(tong)


async def test_lo_het_ton_thuoc_da_tat_van_tra_duoc_tim_theo_ma_tem(co: Co) -> None:
    drug = await _thuoc_moi(co.q)
    lo = await _nhap_kho(co, drug, 3)
    await PharmacyService(co.q.pool).huy(
        identity=co.ds, drug_batch_id=lo, so_luong=3, ly_do="Vỡ cả hộp"
    )
    await co.q.pool.execute(
        "UPDATE drug_catalog SET is_active = false WHERE id = $1::uuid", drug
    )
    ma_lo, ma_tem = await co.q.pool.fetchrow(
        "SELECT batch_code, ma_tem FROM drug_batch WHERE id = $1::uuid", lo
    )

    kq = await _xem(co, drug)
    assert kq["thuoc"]["dang_dung"] is False
    [x] = kq["dang_o_dau"]["lo"]
    assert (x["so_lo"], x["o_quay"], x["da_huy"]) == (ma_lo, 0, 3)
    assert {d["ma"] for d in kq["dong"]} == {"RECEIVE", "DISCARD"}

    for chu in (ma_lo, ma_tem, ma_tem.lower()):
        [g] = await td.tim_thuoc(co.q.pool, identity=co.ds, q=chu)
        assert (g["id"], g["so_lo"], g["dang_dung"]) == (drug, ma_lo, False)


async def test_dong_mo_coi_va_ly_do_ky_thuat_van_hien_chu_de_doc(co: Co) -> None:
    drug = await _thuoc_moi(co.q)
    ma = "3f2b8c1e-6b1a-4f6e-9a55-0c1d2e3f4a5b"
    lo = await _lo_tay(co, drug, "viên", 20, f"nhan-ban-co-so HN từ lô KN {ma}")
    await _ghi_so(co, lo, "ADJUST", -1, f"Đếm lại {ma} cuối ca")
    await _ghi_so(co, lo, "ADJUST", -1, "seed-demo-lo-1")
    # Giao không còn phiếu: đơn gốc không tồn tại.
    await _ghi_so(co, lo, "DISPENSE", -2, None, ref=str(uuid.uuid4()))

    giao, ma_may, co_ma, ton_dau = await _tat_ca(co, drug)
    assert ton_dau["ma"] == "TON_DAU" and ton_dau["hanh_dong"] == "Tồn đầu khi mở cơ sở"
    assert (ton_dau["tu"], ton_dau["ly_do"]) == ("Cơ sở khác", None)
    assert co_ma["ly_do"] == "Đếm lại cuối ca"
    assert ma_may["ly_do"] is None and ma_may["hanh_dong"] == "Điều chỉnh"
    assert giao["hanh_dong"] == "Giao cho khách" and giao["ton_sau"] == 16
    assert giao["khach"] == "Không rõ khách"
    assert (giao["ma_bn"], giao["visit_id"], giao["ma_phieu"]) == (None, None, None)

    hien = [
        str(d[k])
        for d in (giao, ma_may, co_ma, ton_dau)
        for k in ("hanh_dong", "tu", "den", "ly_do", "ma_phieu", "khach")
    ]
    assert not any("nhan-ban" in x or ma in x or "seed" in x for x in hien)


async def test_don_da_dinh_chinh_van_ra_khach(co: Co) -> None:
    rx, drug = await _dong_da_xac_dinh(co.q, 2)
    lo = await _nhap_kho(co, drug, 5)
    await _thu(co.q)
    await PharmacyService(co.q.pool).cap_phat(
        identity=co.ds, prescription_id=rx, drug_batch_id=lo, so_luong=2
    )
    async with co.q.pool.acquire() as conn, conn.transaction():
        lan = await _lan(conn, co.q)
        await conn.execute(
            "UPDATE prescription SET removed_in_correction_id = $2::uuid"
            " WHERE id = $1::uuid",
            rx,
            lan,
        )
    giao = (await _xem(co, drug, hanh_dong="DISPENSE"))["dong"][0]
    assert giao["khach"] and giao["khach"] != "Không rõ khách"
    assert giao["ma_bn"] and giao["visit_id"] == co.q.visit_id


async def test_excel_va_duong_http_bam_bo_loc_rac_khong_500(co: Co) -> None:
    v = await _vong_doi(co)
    app.dependency_overrides[get_db_pool] = lambda: co.q.pool
    app.dependency_overrides[_resolve_identity] = lambda: co.ds
    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://kho"
        ) as c:
            url = "/api/v1/pharmacy/theo-doi-thuoc"
            ra = await c.get(
                f"{url}.xlsx", params={"thuoc": v.drug, "hanh_dong": "DISPENSE"}
            )
            assert ra.status_code == 200
            so = openpyxl.load_workbook(io.BytesIO(ra.content))["Dòng thời gian"]
            giao = await _tat_ca(co, v.drug, hanh_dong="DISPENSE")
            assert so.max_row == len(giao) + 1
            assert {r[2].value for r in so.iter_rows(min_row=2)} == {"Giao cho khách"}

            for tham_so in (
                {},
                {"thuoc": "abc"},
                {"thuoc": v.drug, "tu": "rác", "den": "2026-99-99", "trang": "x"},
                {"thuoc": v.drug, "tu": "9999-12-31", "den": "2026-01-01"},
                {"thuoc": v.drug, "noi": "'; DROP TABLE x; --", "nguoi": "abc"},
            ):
                assert (await c.get(url, params=tham_so)).status_code == 200
                assert (await c.get(f"{url}.xlsx", params=tham_so)).status_code == 200
            assert (await c.get(f"{url}/tim", params={"q": ""})).json() == {"items": []}
            kq = (await c.get(url, params={"thuoc": v.drug})).json()
            # Không lộ SĐT, kể cả trong tệp đã biết tên khách.
            assert not any("sdt" in k or "phone" in k for k in kq["dong"][0])
    finally:
        app.dependency_overrides.pop(get_db_pool, None)
        app.dependency_overrides.pop(_resolve_identity, None)


def test_quyen_doc_dung_bang_the_kho() -> None:
    def quyen(duong: str) -> list[Any]:
        route = next(r for r in pharmacy_router.routes if r.path == duong)  # type: ignore[attr-defined]
        return [
            d.call.quyen
            for d in route.dependant.dependencies  # type: ignore[attr-defined]
            if hasattr(d.call, "quyen")
        ]

    goc = quyen("/pharmacy/the-kho/{drug_catalog_id}")
    assert goc == [("pharmacy.view", "pharmacy.dispense")]
    for duong in (
        "/pharmacy/theo-doi-thuoc",
        "/pharmacy/theo-doi-thuoc.xlsx",
        "/pharmacy/theo-doi-thuoc/tim",
    ):
        assert quyen(duong) == goc


async def test_co_so_khac_khong_lot(co: Co) -> None:
    v = await _vong_doi(co)
    # Cùng thuốc (danh mục chung) có lô ở cơ sở khác.
    await _nhap_lo(co.q, v.drug, 7)

    ben_kia = await td.theo_doi_thuoc(co.q.pool, identity=co.q.duoc_si, thuoc=v.drug)
    [t] = ben_kia["dang_o_dau"]["tong"]
    assert (t["tong"], t["cho_gan_lo"], t["cho_kiem"], t["da_huy"]) == (7, 0, 0, 0)
    assert {d["ma"] for d in ben_kia["dong"]} == {"RECEIVE"}

    ben_nay = await _xem(co, v.drug)
    assert all(d["so_lo"] != ben_kia["dong"][0]["so_lo"] for d in ben_nay["dong"])
    assert 7 not in [x["o_quay"] for x in ben_nay["dang_o_dau"]["lo"]]

    # Tìm số lô chỉ ra lô ở cơ sở đang làm.
    tim = await td.tim_thuoc(co.q.pool, identity=co.q.duoc_si, q=v.ma_lo)
    assert tim == []

    # Phòng khám khác không đọc được thuốc này.
    la = dataclasses.replace(co.ds, clinic_id=str(uuid.uuid4()))
    assert (await td.theo_doi_thuoc(co.q.pool, identity=la, thuoc=v.drug))[
        "thuoc"
    ] is None


async def test_tim_theo_ten_khong_dau_ma_hang_va_so_lo(co: Co) -> None:
    drug = await _thuoc_moi(co.q)
    ten = await co.q.pool.fetchval(
        "UPDATE drug_catalog SET name_raw = 'Đường huyết ' || name_raw,"
        " ma_hang = $2 WHERE id = $1::uuid RETURNING name_raw",
        drug,
        f"MH-{uuid.uuid4().hex[:8]}",
    )
    ma_lo = f"Lo-{uuid.uuid4().hex[:10]}"
    await _nhap_kho(co, drug, 3, ma=ma_lo)

    async def tim(chu: str) -> list[dict[str, Any]]:
        return await td.tim_thuoc(co.q.pool, identity=co.ds, q=chu)

    assert drug in [d["id"] for d in await tim("duong huyet")]
    assert [d["id"] for d in await tim(ten)] == [drug]
    [theo_lo] = await tim(ma_lo.upper())
    assert (theo_lo["id"], theo_lo["so_lo"]) == (drug, ma_lo)


@pytest.mark.parametrize(
    "rac", ["", "   ", "abc", "1' OR '1'='1", "%", "_", "\\", "2026-13-45", "😀" * 200]
)
async def test_dau_vao_rac_tra_rong_khong_nem(q: Quay, rac: str) -> None:
    rx, drug = await _dong_da_xac_dinh(q, 1)
    await _nhap_lo(q, drug, 5)
    ds = q.duoc_si
    kq = await td.theo_doi_thuoc(q.pool, identity=ds, thuoc=rac)
    assert kq["thuoc"] is None and kq["dong"] == [] and kq["tong"] == 0
    assert kq["dang_o_dau"] == {"tong": [], "lo": []}

    # Ngày / trang rác: bỏ qua, vẫn đọc được.
    tot = await td.theo_doi_thuoc(
        q.pool, identity=ds, thuoc=drug, tu=rac, den=rac, trang=rac
    )
    assert tot["tong"] == 1 and tot["trang"] == 1

    # Ô lọc rác: không khớp gì (ô trống thì như không lọc), không ném.
    for o in ("hanh_dong", "noi", "lo", "nguoi"):
        kw: dict[str, Any] = {o: rac}
        ra = await td.theo_doi_thuoc(q.pool, identity=ds, thuoc=drug, **kw)
        assert ra["tong"] == (1 if not rac.strip() else 0)

    # Không có wildcard: "%" không ra cả kho.
    assert isinstance(await td.tim_thuoc(q.pool, identity=ds, q=rac), list)
    assert await td.tim_thuoc(q.pool, identity=ds, q="%") == []


async def test_thuoc_chua_co_lo_van_doc_duoc(q: Quay) -> None:
    drug = await _thuoc_moi(q)
    kq = await td.theo_doi_thuoc(q.pool, identity=q.duoc_si, thuoc=drug)
    assert kq["thuoc"]["id"] == drug
    assert kq["dang_o_dau"]["tong"] == [] and kq["dong"] == [] and kq["tong"] == 0


# ── Hàm thuần: bảng nhãn ─────────────────────────────────────────────────


def test_loai_la_van_co_ten_de_hieu() -> None:
    assert td.nhan_hanh_dong("LOAI_MOI_CHUA_KHAI") == td.HANH_DONG_LA
    assert td.nhan_hanh_dong(None) == td.HANH_DONG_LA
    assert td.nhan_hanh_dong("DISPENSE_REVERSAL") == "Hoàn tác giao thuốc"
    assert td.tu_den("LOAI_MOI_CHUA_KHAI", "QUAY") == ("Quầy", None)
    assert td.tu_den("DISPENSE_REVERSAL", "QUAY") == ("Khách", "Quầy")
    assert td.tu_den("HOAN_TAC_CHUYEN", "KHO", huong="QUAY_KHO") == ("Quầy", "Kho")
    assert td.tu_den("DISCARD", "QUAY", ref_type="drug_return") == (
        "Chờ kiểm",
        "Đã huỷ",
    )
    assert td.tu_den("ADJUST", None) == ("—", None)


def test_ly_do_de_doc() -> None:
    assert td.ly_do_de_doc("  Vỡ   hộp ") == "Vỡ hộp"
    assert td.ly_do_de_doc(None) is None
    assert td.ly_do_de_doc("nhan-ban-co-so HN từ lô KN abc") is None
    assert td.ly_do_de_doc("nhap-tu-script") is None
    assert td.ly_do_de_doc("3f2b8c1e-6b1a-4f6e-9a55-0c1d2e3f4a5b") is None
    # Lời người ghi có mã phiếu vẫn giữ.
    assert td.ly_do_de_doc("Phiếu nhập PN-HN-0001") == "Phiếu nhập PN-HN-0001"


def test_khoang_loc_khong_bao_gio_nem() -> None:
    assert td.khoang_loc("rác", None) == (None, None)
    assert td.khoang_loc("2026-02-30", "2026-13-01") == (None, None)
    a, b = td.khoang_loc("2026-10-31", "2026-10-01")
    assert a is not None and b is not None and a < b
    assert td.khoang_loc("9999-12-31", "9999-12-31")[0] is not None


def test_gop_theo_don_vi_khong_cong_hop_voi_vien() -> None:
    lo = [
        {"don_vi": "Viên", "o_kho": 5, "o_quay": 1, "cho_kiem": 0, "da_huy": 2},
        {"don_vi": "viên ", "o_kho": 0, "o_quay": "2.5", "cho_kiem": 1, "da_huy": 0},
        {"don_vi": "hộp", "o_kho": 3, "o_quay": 0, "cho_kiem": 0, "da_huy": 0},
    ]
    ra = td.gop_theo_don_vi(lo, [{"don_vi": "viên", "so_luong": 2}])
    assert [x["don_vi"] for x in ra] == ["Viên", "hộp"]
    vien = ra[0]
    assert vien["tong"] == D("8.5") and vien["cho_gan_lo"] == 2 and vien["da_huy"] == 2
    assert vien["tren_ke"] == D("6.5")
    assert ra[1]["tong"] == 3 and ra[1]["cho_gan_lo"] == 0
