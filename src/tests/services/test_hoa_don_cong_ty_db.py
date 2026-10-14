"""Hoá đơn công ty (G5c) trên Postgres thật: bản chụp không đổi theo hồ sơ khách,
lịch sử sửa, MST chỉ nhắc, quyền theo lego, file Excel MISA, bấm ngẫu nhiên."""

from __future__ import annotations

import io
import random
from dataclasses import replace
from typing import Any
from uuid import uuid4

import asyncpg
import pytest
from openpyxl import load_workbook

from clinicai.api.identity import ClinicRole, StaffIdentity
from clinicai.core.clock import now_vn
from clinicai.core.exceptions import (
    ResourceNotFoundError,
    SafetyGateError,
    ValidationError,
)
from clinicai.services.hoa_don_cong_ty_service import HoaDonCongTyService
from clinicai.services.quay_thu_service import ma_phieu
from clinicai.services.xuat_hoa_don_service import (
    DONG_DAU,
    XuatHoaDonService,
    xlsx_misa,
)
from tests.services.test_quyen_lego_man_db import _nguoi as nguoi_theo_lego
from tests.services.test_tien_thuoc_cp1_db import tao_quay

pytest_plugins = ["tests.services.test_luot_kham_service_db"]
pytestmark = [pytest.mark.db, pytest.mark.asyncio]

GOI_KE_TOAN = ["bao_cao", "ke_toan", "ds_benh_nhan", "lich_lam_viec"]


async def _phieu(
    pool: asyncpg.Pool, q: Any, manager: StaffIdentity, tien: int = 300_000
) -> str:
    """Một phiếu thu dịch vụ đã thu hôm nay, một dòng. Trả id phiếu."""
    cycle = str(uuid4())
    await pool.execute(
        "INSERT INTO payment_cycle(payment_cycle_id,clinic_id,visit_id,kind,"
        "amount,status,legacy,paid_at,created_at,created_by,confirmed_by,method)"
        " VALUES($1::uuid,$2::uuid,$3::uuid,'dich_vu',$5,'PAID',true,now(),"
        "now(),$4::uuid,$4::uuid,'CASH')",
        cycle,
        manager.clinic_id,
        q.visit_id,
        manager.staff_id,
        tien,
    )
    await pool.execute(
        "INSERT INTO payment_bill_line(clinic_id,payment_cycle_id,visit_id,kind,"
        "source_type,source_id,name_snapshot,quantity,unit,unit_price,line_total,"
        "billing_owner) VALUES($1::uuid,$2::uuid,$3::uuid,'dich_vu','exam',$4,"
        "'Khám phụ khoa',1,NULL,$5,$5,'CLINIC')",
        manager.clinic_id,
        cycle,
        q.visit_id,
        f"x-{uuid4().hex[:6]}",
        tien,
    )
    return cycle


async def _khach(pool: asyncpg.Pool, q: Any) -> str:
    """Khách của lượt thử + thông tin CÁ NHÂN (kiểm không lọt ra HĐ công ty)."""
    pid = await pool.fetchval(
        "UPDATE patient p SET address=NULL, address_detail='12 Kim Ngưu',"
        " ward_name='Phường Vĩnh Tuy', province_name='Hà Nội',"
        " national_id_number='001190000001', email='ca-nhan@example.com'"
        " FROM visit v WHERE v.visit_id=$1::uuid"
        " AND p.clinic_patient_id=v.clinic_patient_id"
        " RETURNING p.clinic_patient_id::text",
        q.visit_id,
    )
    return str(pid)


async def _xuat(
    pool: asyncpg.Pool, manager: StaffIdentity, cycles: list[str]
) -> dict[str, dict[str, Any]]:
    day = now_vn().date().isoformat()
    data = await XuatHoaDonService(pool).xem(identity=manager, tu=day, den=day)
    theo_ma = {h["ma_phieu"]: h for h in data["hoa_don"]}
    return {c: theo_ma[ma_phieu(c)] for c in cycles}


def _dong_xlsx(hd: dict[str, Any]) -> list[Any]:
    ws: Any = load_workbook(io.BytesIO(xlsx_misa([hd]))).active
    return [ws.cell(DONG_DAU, c).value for c in range(1, 18)]


CT_A: dict[str, Any] = {
    "ten_cong_ty": "Công ty TNHH Hoa Sen",
    "dia_chi_cong_ty": "99 Trần Duy Hưng, Hà Nội",
    "ma_so_thue": "0123456789",
    "email_nhan_hd": "ketoan@hoasen.vn",
}


async def test_bat_bien_ban_chup_khong_doi_theo_ho_so_va_xuat_dung_cot(
    pool: asyncpg.Pool,
) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    pid = await _khach(pool, q)
    c1, c2 = await _phieu(pool, q, manager), await _phieu(pool, q, manager)
    svc = HoaDonCongTyService(pool)

    ghi = await svc.ghi_hoa_don(identity=manager, payment_cycle_id=c1, **CT_A)
    assert ghi["canh_bao_mst"] is None and ghi["doi"] is True
    khach_ct_id = ghi["khach_cong_ty_id"]
    # Công ty gõ ở hoá đơn cũng vào HỒ SƠ khách (điền sẵn lần sau).
    book = await svc.doc_cong_ty_khach(identity=manager, clinic_patient_id=pid)
    assert [b["id"] for b in book] == [khach_ct_id]

    # Sửa công ty TRONG HỒ SƠ → bản chụp trên hoá đơn đã ghi KHÔNG đổi.
    await svc.luu_cong_ty_khach(
        identity=manager,
        clinic_patient_id=pid,
        id_=khach_ct_id,
        ten_cong_ty="Công ty Hoa Sen ĐỔI TÊN",
        dia_chi_cong_ty="Địa chỉ mới",
        ma_so_thue="9999999999",
    )
    snap = (await svc.doc_hoa_don(identity=manager, payment_cycle_id=c1))["cong_ty"]
    assert snap["ten_cong_ty"] == CT_A["ten_cong_ty"]
    assert snap["ma_so_thue"] == CT_A["ma_so_thue"]
    assert await svc.doc_lich_su(identity=manager, payment_cycle_id=c1) != []
    assert len(await svc.doc_lich_su(identity=manager, payment_cycle_id=c1)) == 1

    hd = (await _xuat(pool, manager, [c1]))[c1]
    cot = _dong_xlsx(hd)
    # C tên cty · D địa chỉ cty · E MST · F người mua TRỐNG · G email cty · H/I TRỐNG.
    assert cot[2:9] == [
        CT_A["ten_cong_ty"],
        CT_A["dia_chi_cong_ty"],
        CT_A["ma_so_thue"],
        None,
        CT_A["email_nhan_hd"],
        None,
        None,
    ]
    # Không thông tin cá nhân nào trong cả dòng.
    assert "001190000001" not in repr(cot) and "Kim Ngưu" not in repr(cot)
    assert "ca-nhan@example.com" not in repr(cot)

    # Hoá đơn khác (chưa chọn công ty) GIỮ NGUYÊN như trước G5c.
    hd2 = (await _xuat(pool, manager, [c2]))[c2]
    assert not hd2["la_hd_cong_ty"]
    assert hd2["cccd"] == "001190000001" and hd2["email"] == "ca-nhan@example.com"
    assert hd2["dia_chi"] == "12 Kim Ngưu, Phường Vĩnh Tuy, Hà Nội"
    # Chọn công ty không đổi một đồng.
    assert hd["tong"] == hd2["tong"] == 300_000


async def test_sua_ban_chup_ghi_lich_su_bo_va_ghi_lai(pool: asyncpg.Pool) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    await _khach(pool, q)
    c = await _phieu(pool, q, manager)
    svc = HoaDonCongTyService(pool)

    await svc.ghi_hoa_don(identity=manager, payment_cycle_id=c, **CT_A)
    # Ghi lại Y NGUYÊN → không thêm lịch sử.
    lai = await svc.ghi_hoa_don(identity=manager, payment_cycle_id=c, **CT_A)
    assert lai["doi"] is False
    sua = await svc.ghi_hoa_don(
        identity=manager,
        payment_cycle_id=c,
        **{**CT_A, "ma_so_thue": "0123456789-001", "nguoi_mua": "Trần Văn B"},
    )
    assert sua["doi"] is True
    ls = await svc.doc_lich_su(identity=manager, payment_cycle_id=c)
    assert [x["hanh_dong"] for x in ls] == ["sua", "ghi"]
    assert (ls[0]["cu_ma_so_thue"], ls[0]["moi_ma_so_thue"]) == (
        "0123456789",
        "0123456789-001",
    )
    assert ls[0]["moi_nguoi_mua"] == "Trần Văn B" and ls[0]["sua_boi_ten"]

    # Bỏ chọn: hoá đơn trở lại cá nhân, KHÔNG xoá dòng nào; hoàn tác bằng ghi lại.
    await svc.bo_hoa_don(identity=manager, payment_cycle_id=c)
    doc = await svc.doc_hoa_don(identity=manager, payment_cycle_id=c)
    assert doc["cong_ty"] is None and doc["da_bo"] is True
    # Công ty vừa bỏ vẫn đọc được để "Chọn lại" điền sẵn đúng nó (hoàn tác).
    assert doc["cong_ty_da_bo"]["ten_cong_ty"] == CT_A["ten_cong_ty"]
    assert not (await _xuat(pool, manager, [c]))[c]["la_hd_cong_ty"]
    # Bỏ lần hai: im lặng, không thêm lịch sử.
    await svc.bo_hoa_don(identity=manager, payment_cycle_id=c)
    lai2 = await svc.ghi_hoa_don(identity=manager, payment_cycle_id=c, **CT_A)
    assert lai2["doi"] is True
    ls = await svc.doc_lich_su(identity=manager, payment_cycle_id=c)
    assert [x["hanh_dong"] for x in ls] == ["lai", "bo", "sua", "ghi"]
    assert (await _xuat(pool, manager, [c]))[c]["la_hd_cong_ty"]


async def test_mst_sai_van_luu_chi_nhac_va_dau_vao_rac(pool: asyncpg.Pool) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    pid = await _khach(pool, q)
    c = await _phieu(pool, q, manager)
    svc = HoaDonCongTyService(pool)

    r = await svc.ghi_hoa_don(
        identity=manager,
        payment_cycle_id=c,
        **{**CT_A, "ma_so_thue": "12345"},
    )
    assert r["canh_bao_mst"] and "kiểm lại" in r["canh_bao_mst"]
    snap = (await svc.doc_hoa_don(identity=manager, payment_cycle_id=c))["cong_ty"]
    assert snap["ma_so_thue"] == "12345"  # sai dạng vẫn LƯU như người gõ
    # MST có khoảng trắng → chuẩn hoá, không nhắc.
    r = await svc.ghi_hoa_don(
        identity=manager,
        payment_cycle_id=c,
        **{**CT_A, "ma_so_thue": "0123 456 789"},
    )
    assert r["canh_bao_mst"] is None

    # GHI rác → 422, không ghi nửa chừng.
    for rac in ("", "   ", "A" * 401):
        with pytest.raises(ValidationError):
            await svc.ghi_hoa_don(
                identity=manager,
                payment_cycle_id=c,
                **{**CT_A, "ten_cong_ty": rac},
            )
    with pytest.raises(ValidationError):
        await svc.ghi_hoa_don(
            identity=manager, payment_cycle_id="không-phải-uuid", **CT_A
        )
    with pytest.raises(ValidationError):
        await svc.ghi_hoa_don(
            identity=manager,
            payment_cycle_id=c,
            khach_cong_ty_id="rác",
            **CT_A,
        )
    with pytest.raises(ValidationError):
        await svc.luu_cong_ty_khach(
            identity=manager, clinic_patient_id="rác", ten_cong_ty="A"
        )
    with pytest.raises(ResourceNotFoundError):
        await svc.ghi_hoa_don(identity=manager, payment_cycle_id=str(uuid4()), **CT_A)
    with pytest.raises(ResourceNotFoundError):
        await svc.luu_cong_ty_khach(
            identity=manager, clinic_patient_id=str(uuid4()), ten_cong_ty="A"
        )
    with pytest.raises(ResourceNotFoundError):
        await svc.bo_hoa_don(identity=manager, payment_cycle_id=str(uuid4()))
    # ĐỌC rác → rỗng, không ném.
    assert await svc.doc_cong_ty_khach(identity=manager, clinic_patient_id="rác") == []
    assert (await svc.doc_hoa_don(identity=manager, payment_cycle_id="rác"))[
        "cong_ty"
    ] is None
    assert await svc.doc_lich_su(identity=manager, payment_cycle_id="rác") == []
    # Sau tất cả bài rác: bản chụp vẫn là lần ghi hợp lệ cuối.
    snap = (await svc.doc_hoa_don(identity=manager, payment_cycle_id=c))["cong_ty"]
    assert snap["ten_cong_ty"] == CT_A["ten_cong_ty"]
    assert (
        len(await svc.doc_cong_ty_khach(identity=manager, clinic_patient_id=pid)) == 1
    )


async def test_nhieu_cong_ty_dien_san_lan_truoc_va_an_hoan_tac(
    pool: asyncpg.Pool,
) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    pid = await _khach(pool, q)
    c1, c2, c3 = [await _phieu(pool, q, manager) for _ in range(3)]
    svc = HoaDonCongTyService(pool)

    # Lễ tân/CSKH ghi trước vào hồ sơ: hai công ty.
    a = await svc.luu_cong_ty_khach(
        identity=manager, clinic_patient_id=pid, ten_cong_ty="Công ty A"
    )
    b = await svc.luu_cong_ty_khach(
        identity=manager, clinic_patient_id=pid, ten_cong_ty="Công ty B"
    )
    # Chưa dùng lần nào: cái mới tạo đứng đầu.
    goi_y = (await svc.doc_hoa_don(identity=manager, payment_cycle_id=c1))["goi_y"]
    assert [g["id"] for g in goi_y] == [b["id"], a["id"]]
    # Thu ngân chọn công ty A cho hoá đơn 1 → lần sau A đứng đầu (điền sẵn lần trước).
    await svc.ghi_hoa_don(
        identity=manager,
        payment_cycle_id=c1,
        ten_cong_ty="Công ty A",
        khach_cong_ty_id=a["id"],
    )
    goi_y = (await svc.doc_hoa_don(identity=manager, payment_cycle_id=c2))["goi_y"]
    assert [g["id"] for g in goi_y] == [a["id"], b["id"]]
    # Công ty gõ lại y tên (không chọn id) → khớp công ty có sẵn, không nhân đôi.
    await svc.ghi_hoa_don(
        identity=manager, payment_cycle_id=c3, ten_cong_ty="công ty b"
    )
    book = await svc.doc_cong_ty_khach(identity=manager, clinic_patient_id=pid)
    assert len(book) == 2

    # Gỡ khỏi hồ sơ = ẩn (hoàn tác được), không xoá; bản chụp cũ vẫn nguyên.
    await svc.an_cong_ty_khach(identity=manager, id_=a["id"], an=True)
    book = await svc.doc_cong_ty_khach(identity=manager, clinic_patient_id=pid)
    assert [x["id"] for x in book] == [b["id"]]
    tat_ca = await svc.doc_cong_ty_khach(
        identity=manager, clinic_patient_id=pid, ca_da_an=True
    )
    assert {x["id"] for x in tat_ca} == {a["id"], b["id"]}
    snap = (await svc.doc_hoa_don(identity=manager, payment_cycle_id=c1))["cong_ty"]
    assert snap["ten_cong_ty"] == "Công ty A"
    await svc.an_cong_ty_khach(identity=manager, id_=a["id"], an=False)
    book = await svc.doc_cong_ty_khach(identity=manager, clinic_patient_id=pid)
    assert {x["id"] for x in book} == {a["id"], b["id"]}
    # Công ty của khách khác không gắn được vào phiếu này (lọt sang là sai khách).
    q2 = await tao_quay(pool)
    pid2 = await _khach(pool, q2)
    lac = await svc.luu_cong_ty_khach(
        identity=manager, clinic_patient_id=pid2, ten_cong_ty="Công ty lạ"
    )
    r = await svc.ghi_hoa_don(
        identity=manager,
        payment_cycle_id=c2,
        ten_cong_ty="Công ty lạ",
        khach_cong_ty_id=lac["id"],
        luu_vao_ho_so=False,
    )
    assert r["khach_cong_ty_id"] is None


async def test_cong_quyen_theo_lego_khong_theo_vai(pool: asyncpg.Pool) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    pid = await _khach(pool, q)
    c = await _phieu(pool, q, manager)
    svc = HoaDonCongTyService(pool)

    khong = await nguoi_theo_lego(pool, "CSKH", ["bao_cao", "ds_benh_nhan"])
    for goi in (
        svc.doc_hoa_don(identity=khong, payment_cycle_id=c),
        svc.doc_cong_ty_khach(identity=khong, clinic_patient_id=pid),
        svc.ghi_hoa_don(identity=khong, payment_cycle_id=c, **CT_A),
        svc.bo_hoa_don(identity=khong, payment_cycle_id=c),
        svc.luu_cong_ty_khach(identity=khong, clinic_patient_id=pid, ten_cong_ty="A"),
    ):
        with pytest.raises(SafetyGateError):
            await goi
    # Thu ngân (lúc thu) và Kế toán (sửa sau) đều ghi được.
    await svc.ghi_hoa_don(identity=q.thu_ngan, payment_cycle_id=c, **CT_A)
    ke_toan = await nguoi_theo_lego(pool, "CSKH", GOI_KE_TOAN)
    r = await svc.ghi_hoa_don(
        identity=ke_toan,
        payment_cycle_id=c,
        **{**CT_A, "dia_chi_cong_ty": "Kế toán sửa địa chỉ"},
    )
    assert r["doi"] is True
    ls = await svc.doc_lich_su(identity=ke_toan, payment_cycle_id=c)
    assert [x["hanh_dong"] for x in ls] == ["sua", "ghi"]
    assert ls[0]["sua_boi_ten"] != ls[1]["sua_boi_ten"]


async def test_bam_ngau_nhien_40_chuoi_x_10_buoc_khop_mo_hinh(
    pool: asyncpg.Pool,
) -> None:
    """Mô hình trong bộ nhớ so với DB sau MỖI bước: bản chụp, số dòng lịch sử, cột
    Excel, hồ sơ khách không kéo bản chụp theo, và tổng tiền không bao giờ đổi."""
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    pid = await _khach(pool, q)
    svc = HoaDonCongTyService(pool)
    ten_mau = ["Công ty Alpha", "Công ty Beta", "Công ty Gamma", "Hộ KD Delta"]
    mst_mau = ["0123456789", "0123456789-001", "12345", "", "ABCDEFGHIJ"]
    ho_so_ids: list[str] = []
    rng = random.Random(14_10_2026)
    for chuoi in range(40):
        cycles = [
            await _phieu(pool, q, manager, 100_000 + 1000 * chuoi) for _ in range(3)
        ]
        mo_hinh: dict[str, dict[str, Any]] = {
            c: {"cong_ty": None, "bo": False, "ls": 0} for c in cycles
        }
        for _ in range(10):
            c = rng.choice(cycles)
            m = mo_hinh[c]
            op = rng.choice(["ghi", "ghi", "bo", "sua_ho_so", "an_ho_so", "doc_ls"])
            if op == "ghi":
                moi: dict[str, Any] = {
                    "ten_cong_ty": rng.choice(ten_mau),
                    "dia_chi_cong_ty": rng.choice(["", "1 Phố A", "2 Phố B"]),
                    "ma_so_thue": rng.choice(mst_mau),
                    "email_nhan_hd": rng.choice(["", "kt@x.vn"]),
                    "nguoi_mua": rng.choice(["", "Người Mua"]),
                }
                r = await svc.ghi_hoa_don(identity=manager, payment_cycle_id=c, **moi)
                if m["cong_ty"] is None and not m["bo"]:
                    m["ls"] += 1
                elif m["bo"] or m["cong_ty"] != moi:
                    m["ls"] += 1
                else:
                    assert r["doi"] is False
                m["cong_ty"], m["bo"] = moi, False
                if r["khach_cong_ty_id"]:
                    ho_so_ids.append(r["khach_cong_ty_id"])
            elif op == "bo":
                if m["cong_ty"] is None:
                    with pytest.raises(ResourceNotFoundError):
                        await svc.bo_hoa_don(identity=manager, payment_cycle_id=c)
                else:
                    await svc.bo_hoa_don(identity=manager, payment_cycle_id=c)
                    if not m["bo"]:
                        m["ls"] += 1
                    m["bo"] = True
            elif op == "sua_ho_so" and ho_so_ids:
                await svc.luu_cong_ty_khach(
                    identity=manager,
                    clinic_patient_id=pid,
                    id_=rng.choice(ho_so_ids),
                    ten_cong_ty=f"ĐỔI TÊN {rng.random()}",
                    ma_so_thue="9876543210",
                )
            elif op == "an_ho_so" and ho_so_ids:
                await svc.an_cong_ty_khach(
                    identity=manager,
                    id_=rng.choice(ho_so_ids),
                    an=rng.random() < 0.5,
                )
            # Bất biến sau MỖI bước, cho cả 3 phiếu.
            for cy, mm in mo_hinh.items():
                doc = await svc.doc_hoa_don(identity=manager, payment_cycle_id=cy)
                hien = None if mm["bo"] else mm["cong_ty"]
                if hien is None:
                    assert doc["cong_ty"] is None
                else:
                    assert {k: doc["cong_ty"][k] for k in hien} == hien
                ls = await svc.doc_lich_su(identity=manager, payment_cycle_id=cy)
                assert len(ls) == mm["ls"]
        # Cuối chuỗi: file xuất khớp mô hình, tổng tiền không đổi, không lộ cá nhân.
        hds = await _xuat(pool, manager, cycles)
        for cy, mm in mo_hinh.items():
            hd = hds[cy]
            assert hd["tong"] == 100_000 + 1000 * chuoi
            hien = None if mm["bo"] else mm["cong_ty"]
            cot = _dong_xlsx(hd)
            if hien is None:
                assert hd["la_hd_cong_ty"] is False
                assert cot[2] == hd["ten_khach"] and cot[8] == "001190000001"
            else:
                assert cot[2] == hien["ten_cong_ty"]
                assert (cot[3] or "") == hien["dia_chi_cong_ty"]
                assert (cot[4] or "") == hien["ma_so_thue"].replace(" ", "")
                assert cot[7] is None and cot[8] is None
                assert "ca-nhan@example.com" not in repr(cot)
