"""Invoice register isolation and actual refund dates against PostgreSQL."""

from dataclasses import replace
from uuid import uuid4

import asyncpg
import pytest

from clinicai.api.identity import ClinicRole
from clinicai.core.clock import now_vn
from clinicai.core.exceptions import SafetyGateError
from clinicai.services.danh_sach_hoa_don_service import DanhSachHoaDonService
from tests.services.test_tien_thuoc_cp1_db import tao_quay

pytest_plugins = ["tests.services.test_luot_kham_service_db"]
pytestmark = [pytest.mark.db, pytest.mark.asyncio]


async def test_old_paid_bill_refunded_today_and_scope(pool: asyncpg.Pool) -> None:
    q = await tao_quay(pool)
    manager = replace(q.thu_ngan, role=ClinicRole.MANAGEMENT)
    cycle = str(uuid4())
    await pool.execute(
        "INSERT INTO payment_cycle(payment_cycle_id,clinic_id,visit_id,kind,"
        "amount,status,legacy,paid_at,created_at,created_by,confirmed_by)"
        " VALUES($1::uuid,$2::uuid,$3::uuid,'dich_vu',100000,'PAID',true,"
        "now()-interval '2 days',now()-interval '2 days',$4::uuid,$4::uuid)",
        cycle,
        manager.clinic_id,
        q.visit_id,
        manager.staff_id,
    )
    async with pool.acquire() as conn, conn.transaction():
        line = await conn.fetchval(
            "INSERT INTO payment_bill_line(clinic_id,payment_cycle_id,visit_id,"
            "kind,source_type,source_id,name_snapshot,quantity,unit_price,line_total,"
            "billing_owner) VALUES($1::uuid,$2::uuid,$3::uuid,'dich_vu','exam','x',"
            "'Dịch vụ hoàn',1,30000,30000,'CLINIC') RETURNING id",
            manager.clinic_id,
            cycle,
            q.visit_id,
        )
        refund = await conn.fetchval(
            "INSERT INTO payment_refund(clinic_id,visit_id,kind,payment_cycle_id,"
            "amount,status,method,reason,created_by,created_at)"
            " VALUES($1::uuid,$2::uuid,'dich_vu',$3::uuid,30000,'PENDING','CASH',"
            "'khách trả lại',$4::uuid,now()-interval '1 day')"
            " RETURNING refund_id",
            manager.clinic_id,
            q.visit_id,
            cycle,
            manager.staff_id,
        )
        await conn.execute(
            "INSERT INTO payment_refund_line(clinic_id,refund_id,payment_bill_line_id,"
            "payment_cycle_id,quantity,amount) VALUES($1::uuid,$2,$3,$4::uuid,1,30000)",
            manager.clinic_id,
            refund,
            line,
            cycle,
        )
        await conn.execute(
            "UPDATE payment_refund SET status='COMPLETED',completed_at=now(),"
            "completed_by=$2::uuid WHERE refund_id=$1",
            refund,
            manager.staff_id,
        )
    svc = DanhSachHoaDonService(pool)
    day = now_vn().date().isoformat()
    data = await svc.xem(identity=manager, tu=day, den=day, tim=f"TT-{q.duoi}")
    assert data["tong"]["da_thu"] == 0
    assert data["tong"]["da_hoan"] == 30000
    [customer] = data["khach"]
    assert customer["thu_rong"] == 70000
    assert customer["phieu"][0]["nguoi_thu"] == manager.full_name
    assert customer["phieu"][0]["hoan"][0]["nguoi_hoan"] == manager.full_name
    assert not (
        await svc.xem(identity=manager, co_so=str(uuid4()), tim=f"TT-{q.duoi}")
    )["khach"]
    assert not (await svc.xem(identity=manager, co_so="rác", tim=f"TT-{q.duoi}"))[
        "khach"
    ]
    other = str(
        await pool.fetchval(
            "INSERT INTO clinic(code,name,timezone)"
            " VALUES($1,'Khác','Asia/Ho_Chi_Minh')"
            " RETURNING id::text",
            f"HD-{q.duoi}",
        )
    )
    assert not (
        await svc.xem(identity=replace(manager, clinic_id=other), tim=f"TT-{q.duoi}")
    )["khach"]
    with pytest.raises(SafetyGateError):
        await svc.xem(identity=q.thu_ngan, tu=day, den=day)
    assert (await svc.xem(identity=manager, tu="rác", den="rác", tim=f"TT-{q.duoi}"))[
        "tong"
    ]["da_hoan"] == 30000


async def test_head_scope_and_xlsx_endpoint_share_guard(pool: asyncpg.Pool) -> None:
    from datetime import timedelta
    from io import BytesIO

    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient
    from openpyxl import load_workbook

    from clinicai.api.identity import get_current_identity
    from clinicai.api.v1.routers.hoa_don import router
    from clinicai.core.database import get_db_pool

    q = await tao_quay(pool)
    today = now_vn().date()
    day = today.isoformat()
    await pool.execute(
        """INSERT INTO work_roster
        (clinic_id,week_start,work_date,shift,station,staff_id,staff_name,status)
        VALUES($1::uuid,$2::date,$2::date,'SANG','DIEU_PHOI',
               $3::uuid,'Head invoice','APPROVED')""",
        q.thu_ngan.clinic_id,
        today,
        q.thu_ngan.staff_id,
    )
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db_pool] = lambda: pool
    app.dependency_overrides[get_current_identity] = lambda: q.thu_ngan
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        for path in ("/hoa-don", "/hoa-don/xlsx"):
            response = await client.get(
                path,
                params=dict(tu=day, den=day, ca="SANG", co_so=q.thu_ngan.location_id),
            )
            assert response.status_code == 200
            if path.endswith("xlsx"):
                book = load_workbook(BytesIO(response.content))
                assert len(book.sheetnames) == 7  # + Sau khi thu, Việc còn tồn
                assert book["Tổng hợp"]["A2"].is_date
            for params in (
                dict(tu=day, den=day, ca="INVALID"),
                dict(tu=day, den=day, ca="SANG", co_so=str(uuid4())),
                dict(
                    tu=(today - timedelta(days=1)).isoformat(),
                    den=(today - timedelta(days=1)).isoformat(),
                    ca="SANG",
                ),
            ):
                with pytest.raises(SafetyGateError):
                    await client.get(path, params=params)
    await pool.execute(
        "UPDATE work_roster SET station='T1_LETAN' WHERE staff_id=$1::uuid",
        q.thu_ngan.staff_id,
    )
    with pytest.raises(SafetyGateError):
        await DanhSachHoaDonService(pool).xem(
            identity=q.thu_ngan, tu=day, den=day, ca="SANG"
        )
