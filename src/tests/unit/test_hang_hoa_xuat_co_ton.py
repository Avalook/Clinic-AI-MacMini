"""Xuất Excel / CSV báo cáo hàng hoá có cột TỒN KHO (14/10/2026).

Màn + bản in có Tồn đầu · Nhập trong ca · SL bán · Tồn cuối (G6 #445) nhưng tệp xuất
chưa có. Tệp xuất phải đúng thứ tự cột như màn; phần tiền không đổi; nhiều đơn vị
lô thì mỗi đơn vị một dòng, tiền chỉ ở dòng đầu.
"""

from __future__ import annotations

import csv
import io
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from clinicai.services.bao_cao_cuoi_ngay_service import (
    TIEU_DE_CO_TON,
    csv_hang_hoa,
)
from clinicai.services.excel_report_export import xlsx_hang_hoa

#: Thứ tự cột trên màn `BangHangHoa.tsx` (bảng có tồn) — khách chốt.
THU_TU_MAN = [
    "Mã hàng",
    "Tên hàng",
    "Tồn đầu",
    "Nhập trong ca",
    "SL bán",
    "SL trả",
    "Tồn cuối ca",
    "Kiểm kê",
    "Lệch",
    "Đã thu chưa giao",
    "Doanh thu",
    "GT trả",
    "Doanh thu thuần",
]


def _don_vi(dv: str, **kw: Any) -> dict[str, Any]:
    return {
        "don_vi": dv,
        "ton_dau": 10.0,
        "nhap": 5.0,
        "sl_ban": 3.0,
        "sl_tra": 1.0,
        "khac": 0.0,
        "ton_cuoi": 13.0,
        "kiem_ke": None,
        "lech": None,
        "da_thu_chua_giao": 2.0,
    } | kw


def _bc() -> dict[str, Any]:
    thuoc = {
        "co_ton": True,
        "tong": {
            "so_mat_hang": 2,
            "sl_ban": 5.0,
            "doanh_thu": 500_000,
            "sl_tra": 1.0,
            "gia_tri_tra": 20_000,
            "doanh_thu_thuan": 480_000,
        },
        "dong": [
            {
                "ma": "SP1",
                "ten": "Eulac",
                "don_vi": "viên",
                "kho": {
                    "theo_don_vi": [
                        _don_vi("viên"),
                        _don_vi("hộp", ton_dau=1.0, kiem_ke=2.0, lech=0.5),
                    ],
                    "giao_chua_gan_lo": None,
                },
                "sl_ban": 3.0,
                "doanh_thu": 300_000,
                "sl_tra": 1.0,
                "gia_tri_tra": 20_000,
                "doanh_thu_thuan": 280_000,
            },
            {  # không có khối kho: cột tồn trống, SL theo tiền thu
                "ma": "SP2",
                "ten": "Vitamin",
                "don_vi": None,
                "sl_ban": 2.0,
                "doanh_thu": 200_000,
                "sl_tra": 0.0,
                "gia_tri_tra": 0,
                "doanh_thu_thuan": 200_000,
            },
        ],
    }
    return {
        "tu": "2026-10-14",
        "den": "2026-10-14",
        "ca": None,
        "hang_hoa": {"thuoc": thuoc, "dich_vu": {"tong": {}, "dong": []}},
    }


def test_thu_tu_cot_giong_man() -> None:
    assert TIEU_DE_CO_TON == THU_TU_MAN


def test_csv_co_ton_dung_thu_tu_va_tien_khong_nhan_doi() -> None:
    noi = csv_hang_hoa(_bc(), "thuoc")
    rows = list(csv.reader(io.StringIO(noi.lstrip("﻿"))))
    i = rows.index(THU_TU_MAN)
    tong, vien, hop, vit = rows[i + 1 : i + 5]
    assert tong[1] == "SL mặt hàng: 2" and tong[10:] == ["500000", "20000", "480000"]
    # Nhiều đơn vị: mỗi đơn vị một dòng; mã + tiền chỉ ở dòng đầu.
    assert vien[:2] == ["SP1", "Eulac — viên"]
    assert vien[2:7] == ["10", "5", "3", "1", "13"]
    assert vien[10:] == ["300000", "20000", "280000"]
    assert hop[0] == "" and hop[1].endswith("— hộp") and hop[2] == "1"
    assert hop[7:9] == ["2", "0.5"] and hop[10:] == ["", "", ""]
    # Chưa kiểm kê → ô trống (không suy ra).
    assert vien[7] == "" and vien[8] == ""
    # Không có khối kho → cột tồn trống, SL theo tiền thu như màn.
    assert vit[:2] == ["SP2", "Vitamin"] and vit[2:4] == ["", ""]
    assert vit[4:6] == ["2", "0"] and vit[10:] == ["200000", "0", "200000"]


def test_xlsx_co_ton_so_that_dung_thu_tu() -> None:
    ws = load_workbook(BytesIO(xlsx_hang_hoa(_bc(), "thuoc"))).active
    dong = list(ws.values)
    assert list(dong[0]) == THU_TU_MAN
    assert dong[1][2] == 10 and dong[1][6] == 13 and dong[1][10] == 300_000
    assert dong[2][10] is None, "đơn vị thứ hai không lặp tiền"
    assert dong[3][2] is None and dong[3][4] == 2


def test_khong_co_ton_giu_nguyen_khuon_cu() -> None:
    bc = _bc()
    bc["hang_hoa"]["thuoc"]["co_ton"] = False
    noi = csv_hang_hoa(bc, "thuoc")
    assert "Tồn đầu" not in noi and "Doanh thu thuần" in noi
    ws = load_workbook(BytesIO(xlsx_hang_hoa(bc, "thuoc"))).active
    assert list(next(iter(ws.values)))[2] == "Đơn vị"
