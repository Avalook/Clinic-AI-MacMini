"""Money history must distinguish actual refunds from voids."""

from datetime import datetime, timezone
from typing import Any

from clinicai.services.danh_sach_hoa_don_service import gom_hoa_don


def _o(ws: Any, nhan: str) -> Any:
    """Ô dòng 2 dưới tiêu đề ``nhan`` — tra theo TÊN cột, không theo chữ cái."""
    cot = {c.value: c.column for c in ws[1]}
    return ws.cell(2, cot[nhan]).value


def test_void_does_not_invent_refund_and_groups_by_patient() -> None:
    now = datetime(2026, 10, 13, tzinfo=timezone.utc)
    receipts = [
        dict(
            id="a",
            patient_id="p",
            ten="Khách",
            ma_bn="P1",
            amount=100,
            status="VOIDED",
            paid_at=now,
            trong_ky=True,
            huy_trong_ky=True,
        ),
        dict(
            id="b",
            patient_id="p",
            ten="Khách",
            ma_bn="P1",
            amount=50,
            status="PAID",
            paid_at=now,
            trong_ky=True,
            huy_trong_ky=False,
        ),
    ]
    data = gom_hoa_don(receipts, [], [], [], tim=None, trang_thai=None)
    assert len(data["khach"]) == 1
    assert data["tong"]["da_thu"] == 150
    assert data["tong"]["da_hoan"] == 0
    # Huỷ không phải hoàn (không bịa khoản hoàn) nhưng cũng không còn là tiền
    # của phòng khám: thu ròng = thu − hoàn − huỷ, đúng định nghĩa "thực thu"
    # của báo cáo cuối ngày (staging 12/10 từng lệch 3.255.000đ vì thiếu huỷ).
    assert data["tong"]["gia_tri_huy"] == 100
    assert data["tong"]["thu_rong"] == 50
    t = data["tong"]
    assert t["thu_rong"] == t["da_thu"] - t["da_hoan"] - t["gia_tri_huy"]
    huy = next(p for p in data["khach"][0]["phieu"] if p["trang_thai"] == "DA_HUY")
    assert huy["thu_rong"] == 0 and huy["so_tien"] == 100


def test_completed_date_controls_period_totals_pending_is_separate() -> None:
    receipts = [
        dict(
            id="a",
            patient_id="p",
            amount=100,
            status="PAID",
            trong_ky=False,
            huy_trong_ky=False,
        )
    ]
    refunds = [
        dict(id="r", cycle_id="a", amount=30, status="COMPLETED", trong_ky=True),
        dict(id="s", cycle_id="a", amount=20, status="PENDING", trong_ky=True),
    ]
    data = gom_hoa_don(receipts, refunds, [], [], tim=None, trang_thai=None)
    assert data["tong"]["da_thu"] == 0
    assert data["tong"]["da_hoan"] == 30
    assert data["tong"]["cho_hoan"] == 20
    assert data["tong"]["thu_rong"] == -30
    bill = data["khach"][0]["phieu"][0]
    assert bill["thu_rong"] == 70
    assert bill["trang_thai"] == "CHO_HOAN"
    assert receipts[0] == dict(
        id="a",
        patient_id="p",
        amount=100,
        status="PAID",
        trong_ky=False,
        huy_trong_ky=False,
    )


def test_filter_search_receipt_and_full_customer_totals() -> None:
    receipts = [
        dict(
            id="abcdef1234",
            patient_id="p",
            amount=100,
            status="PAID",
            trong_ky=True,
            huy_trong_ky=False,
            ten="An",
        )
    ]
    assert gom_hoa_don(receipts, [], [], [], tim="PT-ABCDEF12", trang_thai=None)[
        "khach"
    ]
    assert not gom_hoa_don(receipts, [], [], [], tim=None, trang_thai="DA_HUY")["khach"]


def test_zero_bill_is_paid_failed_refund_is_not_money_returned() -> None:
    receipts = [
        dict(
            id="a",
            patient_id="p",
            amount=0,
            status="PAID",
            trong_ky=True,
            huy_trong_ky=False,
        )
    ]
    refunds = [
        dict(
            id="r",
            cycle_id="a",
            amount=20,
            status="FAILED",
            trong_ky=True,
            nguoi_huy="Người đóng",
            closed_reason="Ngân hàng từ chối",
        )
    ]
    data = gom_hoa_don(receipts, refunds, [], [], tim=None, trang_thai=None)
    bill = data["khach"][0]["phieu"][0]
    assert bill["trang_thai"] == "DA_THU"
    assert data["tong"]["da_hoan"] == 0
    assert bill["hoan"][0]["nguoi_huy"] == "Người đóng"


def test_needs_review_filter_and_xlsx_lifetime_period_parity() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    from clinicai.services.danh_sach_hoa_don_service import xlsx_hoa_don

    receipts = [
        dict(
            id="a",
            patient_id="p",
            amount=100,
            status="VOIDED",
            trong_ky=False,
            huy_trong_ky=True,
        ),
        dict(
            id="b",
            patient_id="q",
            amount=50,
            status="PAID",
            trong_ky=True,
            huy_trong_ky=False,
        ),
    ]
    data = gom_hoa_don(
        receipts, [], [], [], tim=None, trang_thai=None, can_kiem_tra=True
    )
    assert len(data["khach"]) == 1
    assert data["khach"][0]["can_doi_soat"] is True
    assert data["khach"][0]["canh_bao"] == ["Phiếu huỷ cần xác minh tiền trả lại"]
    data.update(tu="2026-10-13", den="2026-10-13")
    book = load_workbook(BytesIO(xlsx_hoa_don(data)))
    assert _o(book["Tổng hợp"], "Tiền đã thu") == data["tong"]["da_thu"] == 0
    assert _o(book["Phiếu thu"], "Tiền đã thu") == 100
    # Khách chỉ có một phiếu đã huỷ: thu ròng 0 (tiền không còn ở phòng khám).
    assert _o(book["Theo khách"], "Thực thu") == 0
    assert all(
        not sheet.protection.sheet and not sheet.merged_cells.ranges
        for sheet in book.worksheets
    )


def test_discount_original_actual_and_allocated_are_immutable() -> None:
    from copy import deepcopy
    from decimal import Decimal

    variants = (
        (100, None, 0, None, None),
        (70, 100, 30, "AMOUNT", Decimal("30")),
        (75, 100, 25, "PERCENT", Decimal("25")),
        (0, 100, 100, "PERCENT", Decimal("100")),
    )
    for amount, original, discount, kind, value in variants:
        row = dict(
            id="a",
            patient_id="p",
            amount=amount,
            status="PAID",
            trong_ky=True,
            huy_trong_ky=False,
            original_amount=original,
            discount_amount=discount,
            discount_type=kind,
            discount_value=value,
            discount_reason="Chăm sóc khách" if discount else None,
            nguoi_giam="Thu ngân" if discount else None,
        )
        line = dict(
            id="l",
            cycle_id="a",
            so_luong=1,
            thanh_tien=amount,
            discount_allocated=discount,
        )
        saved = deepcopy((row, line))
        data = gom_hoa_don([row], [], [line], [], tim=None, trang_thai=None)
        bill = data["khach"][0]["phieu"][0]
        assert bill["tong_truoc_giam"] == 100
        assert bill["giam_gia"] == discount
        assert bill["so_tien"] == bill["thu_rong"] == amount
        assert bill["trang_thai"] == "DA_THU"
        assert bill["kieu_giam"] == kind
        assert bill["gia_tri_giam"] == (float(value) if value is not None else None)
        assert bill["nguoi_giam"] == row["nguoi_giam"]
        assert bill["dong"][0]["thanh_tien_truoc_giam"] == 100
        assert bill["dong"][0]["giam_gia"] == discount
        assert data["tong"]["giam_gia"] == data["khach"][0]["giam_gia"] == discount
        assert data["tong"]["tong_truoc_giam"] == 100
        assert (row, line) == saved


def test_discount_excel_columns_and_cross_day_refund_do_not_double_count() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    from clinicai.services.danh_sach_hoa_don_service import xlsx_hoa_don

    row = dict(
        id="a",
        patient_id="p",
        amount=70,
        status="PAID",
        trong_ky=False,
        huy_trong_ky=False,
        original_amount=100,
        discount_amount=30,
        discount_type="AMOUNT",
        discount_value=30,
        discount_reason="Ưu đãi",
        nguoi_giam="Thu ngân",
    )
    line = dict(
        id="l",
        cycle_id="a",
        so_luong=1,
        thanh_tien=70,
        discount_allocated=30,
        ten="Dịch vụ",
        da_bo=False,
    )
    refund = dict(id="r", cycle_id="a", amount=20, status="COMPLETED", trong_ky=True)
    data = gom_hoa_don([row], [refund], [line], [], tim=None, trang_thai=None)
    assert data["tong"]["giam_gia"] == data["tong"]["tong_truoc_giam"] == 0
    assert data["tong"]["thu_rong"] == -20
    assert data["khach"][0]["thu_rong"] == 50
    data.update(tu="2026-10-13", den="2026-10-13")
    book = load_workbook(BytesIO(xlsx_hoa_don(data)))
    for sheet, expected in (
        (
            "Phiếu thu",
            {
                "Tổng tiền hàng": 100,
                "Giảm giá": 30,
                "Lý do giảm": "Ưu đãi",
                "Người giảm": "Thu ngân",
                "Tiền đã thu": 70,
            },
        ),
        ("Dòng phiếu", {"Thành tiền": 70, "Tổng tiền hàng": 100, "Giảm giá": 30}),
        ("Theo khách", {"Giảm giá": 30, "Thực thu": 50}),
        ("Tổng hợp", {"Giảm giá": 0, "Thực thu": -20}),
    ):
        ws = book[sheet]
        cells = {cell.value: ws.cell(2, cell.column) for cell in ws[1]}
        for label, value in expected.items():
            assert cells[label].value == value
            if isinstance(value, int):
                assert cells[label].data_type == "n"


def test_su_kien_sau_thu_theo_gio_va_huy_sau_bao_lau() -> None:
    from datetime import timedelta

    from clinicai.services.hoa_don_projection import gom_hoa_don

    thu = datetime(2026, 10, 13, 6, 0, tzinfo=timezone.utc)
    receipts = [
        dict(
            id="c1",
            patient_id="p",
            visit_id="v",
            kind="dich_vu",
            amount=700,
            status="VOIDED",
            paid_at=thu,
            closed_at=thu + timedelta(seconds=52),
            close_reason="thao tác nhầm",
            nguoi_huy="Lan",
            trong_ky=True,
            huy_trong_ky=True,
        )
    ]
    doi = [
        dict(
            cycle_id="c1",
            method_cu="CASH",
            method_moi="TRANSFER",
            luc=thu + timedelta(seconds=10),
            nguoi="Lan",
            ly_do=None,
        )
    ]
    [khach] = gom_hoa_don(
        receipts, [], [], [], tim=None, trang_thai=None, doi_hinh_thuc=doi
    )["khach"]
    ev = khach["phieu"][0]["su_kien"]
    assert [e["loai"] for e in ev] == ["DOI_HINH_THUC", "HUY_PHIEU"]
    assert ev[0]["mo_ta"] == "Tiền mặt → Chuyển khoản"
    assert ev[1]["sau_giay"] == 52 and ev[1]["so_tien"] == -700
    assert khach["co_phieu_huy"] is True
    assert (khach["da_thu_dich_vu"], khach["da_thu_thuoc"]) == (700, 0)


def test_viec_ton_cu_nhat_truoc_va_khach_con_viec_len_dau() -> None:
    from clinicai.services.hoa_don_projection import gom_hoa_don, gom_viec_ton

    k = dict(patient_id="p2", ten="B", ma_bn="B2")
    ton = gom_viec_ton(
        hoan=[],
        ck=[
            dict(
                k,
                payment_cycle_id="c9",
                visit_id="v2",
                kind="dich_vu",
                amount=150,
                created_at=datetime(2026, 10, 13, 7, tzinfo=timezone.utc),
            )
        ],
        no=[
            dict(
                k,
                visit_id="v3",
                so_tien=400,
                ly_do="quên ví",
                ghi_luc=datetime(2026, 10, 11, 9, tzinfo=timezone.utc),
            )
        ],
        tien_thua={},
        luot={},
    )
    assert [v["loai"] for v in ton] == ["GHI_NO", "CK_CHUA_VE"]
    # Nhóm do máy chủ gắn — link "Nợ" ở báo cáo mở đúng nhóm này ở tab Việc tồn.
    assert [v["nhom_no"] for v in ton] == ["no_khach", "no_khach"]
    now = datetime(2026, 10, 13, tzinfo=timezone.utc)
    receipts = [
        dict(
            id=i,
            patient_id=p,
            visit_id=v,
            kind="thuoc",
            amount=10,
            status="PAID",
            paid_at=now,
            trong_ky=True,
            huy_trong_ky=False,
        )
        for i, p, v in (("a", "p1", "v1"), ("b", "p2", "v2"))
    ]
    data = gom_hoa_don(receipts, [], [], [], tim=None, trang_thai=None, viec_ton=ton)
    assert [g["id"] for g in data["khach"]] == ["p2", "p1"]
    assert data["tong"]["so_viec_ton"] == 2
    assert len(data["khach"][0]["viec_ton"]) == 2
    only = gom_hoa_don(
        receipts, [], [], [], tim=None, trang_thai="CON_VIEC_TON", viec_ton=ton
    )
    assert [g["id"] for g in only["khach"]] == ["p2"]


def _tg(ngay: int, gio: int = 9) -> datetime:
    return datetime(2026, 10, ngay, gio, tzinfo=timezone.utc)


def test_duong_xu_ly_theo_loai_viec_ton() -> None:
    """Nút "Xử lý" (G4a nối màn): mỗi loại việc tồn một đích; quầy thu hôm nay đã
    gom lượt cũ nên chỉ `?luot=`, không truyền ngày; việc chưa có màn riêng thì
    chỉ có hướng dẫn bằng chữ — không link giả."""
    from clinicai.services.hoa_don_projection import duong_xu_ly

    base = dict(visit_id="v1", kind=None, luot_con_mo=False, ngay_thu=None)
    no_dv = duong_xu_ly(dict(base, loai="GHI_NO", kind="dich_vu"))
    assert no_dv is not None
    assert (no_dv["href"], no_dv["lego"]) == (
        "/thu-ngan/dich-vu?luot=v1",
        "payment.service.collect",
    )
    no_thuoc = duong_xu_ly(dict(base, loai="GHI_NO", kind="thuoc"))
    assert no_thuoc is not None and no_thuoc["href"] == "/thu-ngan/thuoc?luot=v1"
    # Chưa rõ quầy → dịch vụ (không để nút vô chủ).
    chua = duong_xu_ly(dict(base, loai="GHI_NO"))
    assert chua is not None and chua["href"] == "/thu-ngan/dich-vu?luot=v1"
    ck = duong_xu_ly(dict(base, loai="CK_CHUA_VE", kind="thuoc"))
    assert ck is not None and ck["nhan"] == "Xác nhận đã nhận tiền"
    assert ck["href"] == "/thu-ngan/thuoc?luot=v1"
    # Phòng khám nợ khách: lượt đã đóng → khối Tiền thừa ở quầy; còn mở → check-out.
    dong = duong_xu_ly(dict(base, loai="CHUA_TRA_LAI"))
    assert dong is not None and dong["href"] == "/thu-ngan/dich-vu?luot=v1"
    # Tiền THUỐC phòng khám nợ khách → hoàn ở quầy THUỐC, lego thu thuốc.
    thuoc = duong_xu_ly(dict(base, loai="CHUA_TRA_LAI", kind="thuoc"))
    assert thuoc is not None and thuoc["href"] == "/thu-ngan/thuoc?luot=v1"
    assert thuoc["nhan"] == "Hoàn tiền thuốc"
    assert thuoc["lego"] == "payment.medicine.collect"
    mo = duong_xu_ly(dict(base, loai="CHUA_TRA_LAI", luot_con_mo=True))
    assert mo is not None and mo["href"] == "/reception/checkout?luot=v1"
    assert mo["lego"] == "reception.checkin.perform"
    # Đã thu mà chưa làm: lượt còn mở → check-out; đã đóng → chỉ hướng dẫn.
    lam_mo = duong_xu_ly(dict(base, loai="CHUA_LAM", luot_con_mo=True))
    assert lam_mo is not None and lam_mo["href"] == "/reception/checkout?luot=v1"
    lam_dong = duong_xu_ly(dict(base, loai="CHUA_LAM"))
    assert lam_dong is not None and "href" not in lam_dong
    assert "không làm" in lam_dong["huong_dan"]
    # Hoàn tiền chưa xong: mở quầy đúng NGÀY THU của phiếu gốc.
    hoan = duong_xu_ly(
        dict(base, loai="HOAN_CHUA_XONG", kind="thuoc", ngay_thu="2026-10-13")
    )
    assert hoan is not None
    assert hoan["href"] == "/thu-ngan/thuoc?ngay=2026-10-13"
    assert duong_xu_ly(dict(base, loai="HOAN_THAT_BAI")) is None  # không biết ngày thu
    # Thiếu lượt / loại lạ → không nút.
    assert duong_xu_ly(dict(base, loai="GHI_NO", visit_id=None)) is None
    assert duong_xu_ly(dict(base, loai="LA_HOAC")) is None


def test_mo_ta_chua_lam_noi_dung_trang_thai_luot() -> None:
    """Bấm thật 14/10: việc tồn nói "lượt đã đóng" trong khi Check-out liệt kê
    chính khách ấy là lượt hôm trước CÒN MỞ. Chữ phải theo trạng thái lượt."""
    from clinicai.services.hoa_don_projection import gom_viec_ton

    dong = [dict(ten="Ghế điện từ trường", so_tien=3_000_000, loai="CHUA_LAM")]
    mo = gom_viec_ton(
        hoan=[],
        ck=[],
        no=[],
        tien_thua={"v1": dict(tong=3_000_000, dong=dong)},
        luot={
            "v1": dict(
                patient_id="p",
                ten="A",
                ma_bn="A1",
                luc=_tg(12),
                luot_tao_luc=_tg(12),
                luot_dong_luc=None,
            )
        },
    )
    assert "lượt hôm trước (12/10) chưa đóng" in mo[0]["mo_ta"]
    assert "đóng lượt hoặc làm tiếp" in mo[0]["mo_ta"]
    assert "đã đóng, khách đã về" not in mo[0]["mo_ta"]
    assert mo[0]["luot_con_mo"] is True and mo[0]["luot_ngay"] == "2026-10-12"
    assert mo[0]["xu_ly"]["href"] == "/reception/checkout?luot=v1"
    da_dong = gom_viec_ton(
        hoan=[],
        ck=[],
        no=[],
        tien_thua={"v1": dict(tong=3_000_000, dong=dong)},
        luot={
            "v1": dict(
                patient_id="p",
                ten="A",
                ma_bn="A1",
                luc=_tg(13),
                luot_tao_luc=_tg(12),
                luot_dong_luc=_tg(13),
            )
        },
    )
    assert "lượt đã đóng, khách đã về" in da_dong[0]["mo_ta"]
    assert (
        da_dong[0]["luot_con_mo"] is False and da_dong[0]["luot_ngay"] == "2026-10-13"
    )


def test_luc_that_dau_ky_lay_ngay_that_cua_luot() -> None:
    """ "1 giờ" gây hiểu lầm: dòng DAU_KY có `luc` = mốc BẬT SỔ, không phải tuổi
    khoản nợ. Hiện ngày thật của lượt (đóng, không thì tạo)."""
    from clinicai.services.danh_sach_hoa_don_service import luc_that, quay_cua_no
    from clinicai.services.so_no_service import NHAN_LOAI

    bat_so = _tg(14, 8)
    dau_ky = dict(
        loai="DAU_KY", luc=bat_so, luot_tao_luc=_tg(10), luot_dong_luc=_tg(11)
    )
    assert luc_that(dau_ky) == _tg(11)
    assert luc_that(dict(dau_ky, luot_dong_luc=None)) == _tg(10)
    assert luc_that(dict(dau_ky, luot_tao_luc=None, luot_dong_luc=None)) == bat_so
    # Dòng sổ thường: giờ sự kiện sinh ra khoản ấy.
    assert luc_that(dict(dau_ky, loai="KHI_VE", luc=_tg(12))) == _tg(12)
    assert NHAN_LOAI["DAU_KY"] == "Nợ từ trước khi bật sổ"
    # Quầy thu khoản nợ: chỉ toàn thuốc → thuốc; còn lại → dịch vụ.
    assert quay_cua_no([dict(loai="prescription"), dict(loai="CK_CHO")]) == "thuoc"
    assert (
        quay_cua_no([dict(loai="prescription"), dict(loai="service_order")])
        == "dich_vu"
    )
    assert quay_cua_no([]) == "dich_vu"


def test_moi_sheet_co_ma_khach_deu_co_ten_khach() -> None:
    """Người đối soát đọc theo TÊN khách — sheet chỉ có mã thì phải tra ngược."""
    from io import BytesIO

    from openpyxl import load_workbook

    from clinicai.services.danh_sach_hoa_don_service import xlsx_hoa_don
    from clinicai.services.hoa_don_projection import gom_hoa_don

    now = datetime(2026, 10, 13, tzinfo=timezone.utc)
    receipts = [
        dict(
            id="c1",
            patient_id="p",
            ten="Nguyễn Thị A",
            ma_bn="BN1",
            visit_id="v",
            kind="dich_vu",
            amount=100,
            status="PAID",
            paid_at=now,
            trong_ky=True,
            huy_trong_ky=False,
        )
    ]
    lines = [
        dict(
            id="l1",
            cycle_id="c1",
            ten="Siêu âm",
            so_luong=1,
            thanh_tien=100,
            da_bo=False,
        )
    ]
    refunds = [
        dict(
            id="r1",
            cycle_id="c1",
            amount=30,
            status="COMPLETED",
            completed_at=now,
            trong_ky=True,
        )
    ]
    data = gom_hoa_don(receipts, refunds, lines, [], tim=None, trang_thai=None)
    data.update(tu="2026-10-13", den="2026-10-13")
    book = load_workbook(BytesIO(xlsx_hoa_don(data)))
    thieu = []
    for ws in book.worksheets:
        dau = [c.value for c in ws[1]]
        if "Mã khách" in dau and "Khách" not in dau:
            thieu.append(ws.title)
    assert not thieu, f"Sheet thiếu tên khách: {thieu}"
    assert _o(book["Dòng phiếu"], "Khách") == "Nguyễn Thị A"
    assert _o(book["Hoàn tiền"], "Khách") == "Nguyễn Thị A"


def test_excel_dung_chu_tren_man_khong_lo_ma_tho() -> None:
    """Tuyền 13/10: file từng ghi "Trong kỳ", "toàn chứng từ", DA_THU, dich_vu,
    CASH… — người đối soát không đối chiếu được với màn."""
    from io import BytesIO

    from openpyxl import load_workbook

    from clinicai.services.danh_sach_hoa_don_service import xlsx_hoa_don
    from clinicai.services.hoa_don_projection import gom_hoa_don

    now = datetime(2026, 10, 13, tzinfo=timezone.utc)
    row = dict(
        id="c1",
        patient_id="p",
        ten="Khách A",
        ma_bn="BN1",
        visit_id="v",
        kind="dich_vu",
        amount=70,
        original_amount=100,
        discount_amount=30,
        discount_type="PERCENT",
        discount_value=30,
        discount_reason="quen",
        method="CASH",
        status="PAID",
        paid_at=now,
        trong_ky=True,
        huy_trong_ky=False,
    )
    line = dict(
        id="l",
        cycle_id="c1",
        ten="Siêu âm",
        so_luong=1,
        thanh_tien=70,
        discount_allocated=30,
        da_bo=False,
    )
    refund = dict(
        id="r",
        cycle_id="c1",
        amount=20,
        status="COMPLETED",
        method="TRANSFER",
        trong_ky=True,
    )
    data = gom_hoa_don([row], [refund], [line], [], tim=None, trang_thai=None)
    data.update(tu="2026-10-13", den="2026-10-13")
    book = load_workbook(BytesIO(xlsx_hoa_don(data)))
    assert book.sheetnames == [
        "Tổng hợp",
        "Theo khách",
        "Phiếu thu",
        "Dòng phiếu",
        "Hoàn tiền",
        "Sau khi thu",
        "Việc còn tồn",
    ]
    assert [c.value for c in book["Theo khách"][1]][:13] == [
        "Khách",
        "Mã khách",
        "Điện thoại",
        "Số phiếu",
        "Tổng tiền hàng",
        "Giảm giá",
        "Tiền đã thu",
        "Tiền đã thu dịch vụ",
        "Tiền đã thu thuốc",
        "Đã trả lại",
        "Chưa trả lại",
        "Thực thu",
        "Trạng thái",
    ]
    tho = {
        "DA_THU",
        "dich_vu",
        "thuoc",
        "CASH",
        "TRANSFER",
        "PERCENT",
        "AMOUNT",
        "COMPLETED",
        "PENDING",
        "HOAN_MOT_PHAN",
    }
    o_tho = [
        (ws.title, c.coordinate, c.value)
        for ws in book.worksheets
        for r in ws.iter_rows()
        for c in r
        if c.value in tho
    ]
    assert not o_tho, f"Còn mã thô trong file: {o_tho}"
    assert not any(
        "toàn chứng từ" in str(c.value) or c.value == "Trong kỳ"
        for ws in book.worksheets
        for c in ws[1]
    )
    assert _o(book["Phiếu thu"], "Hình thức") == "Tiền mặt"
    assert _o(book["Phiếu thu"], "Kiểu giảm") == "%"
    assert _o(book["Hoàn tiền"], "Trạng thái") == "Đã trả lại"
