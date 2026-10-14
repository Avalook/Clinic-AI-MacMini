"""Xuất hoá đơn theo mẫu MISA: hàm dựng hoá đơn (thuần) + file .xlsx từ tệp mẫu."""

import io
import warnings
from typing import Any

from openpyxl import load_workbook

from clinicai.services.xuat_hoa_don_service import (
    DONG_DAU,
    MAU,
    TIEU_DE_MAU,
    dung_hoa_don,
    xlsx_misa,
)

TIEU_DE = [
    "Số thứ tự hóa đơn (*)",
    "Ngày hóa đơn",
    "Tên khách hàng",
    "Địa chỉ cty/cá nhân",
    "Mã số thuế",
    "Người mua hàng",
    "Email",
    "Số điện thoại",
    "CCCD",
    "Hình thức thanh toán",
    "Thuế suất GTGT (%)",
    "Tiền thuế GTGT",
    "Tên hàng hóa/dịch vụ (*)",
    "ĐVT",
    "Số lượng",
    "Đơn giá",
    "Thành tiền",
]


def _phieu(
    id_: str, thu_luc: str, dong: list[dict[str, Any]], **kw: object
) -> dict[str, Any]:
    return {
        "id": id_,
        "ma": f"PT-{id_.upper()}",
        "kind": "dich_vu",
        "thu_luc": thu_luc,
        "thu_trong_ky": True,
        "trang_thai": "DA_THU",
        "hinh_thuc": "TRANSFER",
        "dong": dong,
        "hoan": [],
        **kw,
    }


def _dong(id_: str, ten: str, sl: float, tien: int) -> dict[str, Any]:
    return {"id": id_, "ten": ten, "so_luong": sl, "thanh_tien": tien}


KHACH = [
    {
        "ten": "Nguyễn Thị A",
        "ma_bn": "BN1",
        "so_dien_thoai": "0901",
        "phieu": [
            _phieu(
                "b2",
                "2026-10-12T09:00:00+07:00",
                [_dong("l3", "Thuốc X", 2, 90000)],
                kind="thuoc",
                hinh_thuc="CASH",
            ),
            _phieu(
                "b1",
                "2026-10-12T08:00:00+07:00",
                [
                    _dong("l1", "Khám phụ khoa", 1, 300000),
                    _dong("l2", "=1+1", 1, 50000),
                ],
            ),
            _phieu(
                "b3",
                "2026-10-12T10:00:00+07:00",
                [_dong("l4", "Huỷ", 1, 1)],
                trang_thai="DA_HUY",
            ),
            _phieu(
                "b4",
                "2026-10-01T10:00:00+07:00",
                [_dong("l5", "Ngoài kỳ", 1, 1)],
                thu_trong_ky=False,
            ),
        ],
    }
]
THEM_DONG: dict[str, dict[str, Any]] = {
    "l1": {"unit": None, "source_type": "exam"},
    "l2": {"unit": None, "source_type": "service_order"},
    "l3": {"unit": "Hộp", "source_type": "prescription", "ma": "TX01"},
}
THEM_KHACH: dict[str, dict[str, Any]] = {
    "b1": {"dia_chi": "1 Kim Ngưu, Hà Nội", "cccd": "001", "email": None}
}


def test_moi_phieu_mot_hoa_don_theo_gio_thu_bo_huy_va_ngoai_ky() -> None:
    hd = dung_hoa_don(KHACH, THEM_DONG, THEM_KHACH)
    assert [h["ma_phieu"] for h in hd] == ["PT-B1", "PT-B2"]
    assert [h["stt"] for h in hd] == [1, 2]
    a, b = hd
    assert a["ngay"] == "12/10/2026" and a["hinh_thuc"] == "Chuyển khoản"
    assert a["dia_chi"] == "1 Kim Ngưu, Hà Nội" and a["cccd"] == "001"
    # Hệ thống chưa có → để trống cho kế toán.
    assert a["ma_so_thue"] == "" and a["email"] == ""
    assert [d["dvt"] for d in a["dong"]] == ["Lần", "Lần"]
    assert b["dong"] == [
        {
            "id": "l3",
            "ten": "Thuốc X",
            "ma": "TX01",
            "nhom": "THUOC",
            "dvt": "Hộp",
            "so_luong": 2,
            "don_gia": 45000,
            "thanh_tien": 90000,
            "giam_gia": 0,
            "truoc_giam": 90000,
        }
    ]
    assert b["hinh_thuc"] == "Tiền mặt" and b["dia_chi"] == ""


def test_tru_dong_da_hoan_xong_va_ghi_chu_hoan_cho() -> None:
    p = _phieu(
        "b9",
        "2026-10-12T08:00:00+07:00",
        [_dong("x1", "Siêu âm", 1, 500000), _dong("x2", "Thuốc", 3, 90000)],
        hoan=[
            {
                "trang_thai": "COMPLETED",
                "so_tien": 500000,
                "dong": [{"id": "x1", "thanh_tien": 500000, "so_luong": 1}],
            },
            {
                "trang_thai": "COMPLETED",
                "so_tien": 30000,
                "dong": [{"id": "x2", "thanh_tien": 30000, "so_luong": 1}],
            },
            {"trang_thai": "PENDING", "so_tien": 20000, "dong": []},
        ],
    )
    (hd,) = dung_hoa_don([{"ten": "B", "phieu": [p]}], {}, {})
    assert hd["dong"] == [
        {
            "id": "x2",
            "ten": "Thuốc",
            "ma": "",
            "nhom": "DICH_VU",
            "dvt": "",
            "so_luong": 2,
            "don_gia": 30000,
            "thanh_tien": 60000,
            "giam_gia": 0,
            "truoc_giam": 60000,
        }
    ]
    assert hd["tong"] == 60000
    assert hd["ghi_chu"] == ["Đang chờ hoàn 20.000đ"]


def test_dong_co_giam_gia_ghi_truoc_giam_va_hoan_mot_phan_theo_ti_le() -> None:
    """K2: thanh_tien vẫn là SAU giảm; có thêm số trước giảm cho nút "ghi giảm giá"."""
    p = _phieu(
        "g1",
        "2026-10-12T08:00:00+07:00",
        [
            {**_dong("g1a", "Siêu âm", 1, 270000), "giam_gia": 30000},
            {**_dong("g1b", "Thuốc", 4, 200000), "giam_gia": 40000},
        ],
        hoan=[
            {
                "trang_thai": "COMPLETED",
                "so_tien": 50000,
                "dong": [{"id": "g1b", "thanh_tien": 50000, "so_luong": 1}],
            }
        ],
    )
    (hd,) = dung_hoa_don([{"ten": "G", "phieu": [p]}], {}, {})
    a, b = hd["dong"]
    assert (a["thanh_tien"], a["giam_gia"], a["truoc_giam"]) == (270000, 30000, 300000)
    # Hoàn 1/4 dòng: còn 150.000 sau giảm → giảm còn lại 40.000 × 150/200 = 30.000.
    assert (b["thanh_tien"], b["giam_gia"], b["truoc_giam"]) == (150000, 30000, 180000)
    assert hd["tong"] == 420000


def test_hoan_het_thi_bo_phieu_va_ngay_rac_khong_nem() -> None:
    p = _phieu(
        "b8",
        "rác",
        [_dong("y1", "Khám", 1, 100)],
        hoan=[
            {
                "trang_thai": "COMPLETED",
                "so_tien": 100,
                "dong": [{"id": "y1", "thanh_tien": 100, "so_luong": 1}],
            }
        ],
    )
    assert dung_hoa_don([{"ten": "C", "phieu": [p]}], {}, {}) == []
    q = _phieu("b7", "rác", [_dong("y2", "Khám", 1, 100)])
    assert dung_hoa_don([{"ten": "C", "phieu": [q]}], {}, {})[0]["ngay"] == ""


def test_file_giu_nguyen_mau_va_sua_duoc() -> None:
    hd = dung_hoa_don(KHACH, THEM_DONG, THEM_KHACH)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        mau: Any = load_workbook(MAU).active
        ws: Any = load_workbook(io.BytesIO(xlsx_misa(hd))).active
    assert ws.title == mau.title == "Hóa đơn GTGT"
    # Bỏ 8 dòng chữ hướng dẫn của mẫu: tiêu đề cột lên dòng 1, giữ font, màu,
    # chiều cao, độ rộng cột như mẫu.
    assert DONG_DAU == 2
    assert [ws.cell(1, c).value for c in range(1, 18)] == TIEU_DE
    for c in range(1, 18):
        moi, goc = ws.cell(1, c), mau.cell(TIEU_DE_MAU, c)
        assert (moi.font.name, moi.font.sz, moi.font.b) == (
            goc.font.name,
            goc.font.sz,
            goc.font.b,
        )
        assert repr(moi.fill.fgColor) == repr(goc.fill.fgColor)
        assert moi.border.bottom.style == goc.border.bottom.style
    assert ws.row_dimensions[1].height == mau.row_dimensions[TIEU_DE_MAU].height
    assert ws.column_dimensions["M"].width == mau.column_dimensions["M"].width
    chu = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Hướng dẫn" not in chu and "File mẫu" not in chu
    assert "Nếu là công ty" not in chu
    assert ws.protection.sheet is False
    rows = [
        [ws.cell(r, c).value for c in range(1, 18)]
        for r in range(DONG_DAU, DONG_DAU + 4)
    ]
    # Hoá đơn 1 có 2 dòng: số thứ tự lặp, thông tin khách chỉ ở dòng đầu.
    ten, dia_chi = "Nguyễn Thị A", "1 Kim Ngưu, Hà Nội"
    assert rows[0][:5] == ["1", "12/10/2026", ten, dia_chi, None]
    assert rows[0][5:10] == [ten, None, "0901", "001", "Chuyển khoản"]
    assert rows[0][10:] == [None, None, "Khám phụ khoa", "Lần", 1, 300000, 300000]
    assert rows[1][0] == "1" and rows[1][2] is None
    # Chữ bắt đầu bằng "=" vẫn là chữ, không thành công thức.
    assert rows[1][12] == "=1+1" and ws.cell(DONG_DAU + 1, 13).data_type == "s"
    assert rows[2][0] == "2" and rows[2][12:] == ["Thuốc X", "Hộp", 2, 45000, 90000]
    # Dòng ví dụ + công thức của mẫu đã xoá sạch.
    assert rows[3] == [None] * 17
    assert all(ws.cell(r, 17).value is None for r in range(DONG_DAU + 3, 101))
    # Câu gợi ý khi bấm ô được gắn lại như mẫu.
    assert len(ws.data_validations.dataValidation) == 15


def test_khong_co_hoa_don_van_ra_file_mau_trong() -> None:
    ws: Any = load_workbook(io.BytesIO(xlsx_misa([]))).active
    assert ws.cell(1, 1).value == TIEU_DE[0]
    assert all(c.value is None for row in ws.iter_rows(min_row=2) for c in row)
    assert ws.cell(DONG_DAU, 1).value is None and ws.cell(DONG_DAU, 17).value is None
