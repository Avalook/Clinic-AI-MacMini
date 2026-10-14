"""Thống kê theo hàng hoá (G5a) — gom từ dòng hoá đơn, bảng số tính tay."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from clinicai.services.theo_hang_hoa_service import (
    gom_theo_hang_hoa,
    xlsx_theo_hang_hoa,
)


def _d(ten: str, sl: float, sau: int, giam: int = 0, **x: Any) -> dict[str, Any]:
    return {
        "ten": ten,
        "ma": x.get("ma", ""),
        "nhom": x.get("nhom", "DICH_VU"),
        "dvt": x.get("dvt", "Lần"),
        "so_luong": sl,
        "thanh_tien": sau,
        "giam_gia": giam,
        "truoc_giam": sau + giam,
    }


HOA_DON = [
    {
        "dong": [
            _d("Siêu âm", 1, 270_000, 30_000, ma="SA01"),
            _d("Paracetamol", 10, 50_000, ma="T1", nhom="THUOC", dvt="Viên"),
            _d("Khám phụ khoa", 1, 150_000),
        ]
    },
    {
        "dong": [
            # Cùng mã, tên khác (đổi tên giữa kỳ) → gộp một dòng, giữ tên đầu tiên.
            _d("Siêu âm bụng", 2, 600_000, ma="SA01"),
            # Không mã: so theo tên bỏ khoảng trắng thừa / hoa thường.
            _d("  khám  PHỤ khoa ", 1, 150_000),
            _d("Que thử", 3, 90_000, nhom="VAT_TU", dvt="Cái"),
        ]
    },
]


def test_gom_theo_mat_hang_bang_so_tay() -> None:
    r = gom_theo_hang_hoa(HOA_DON)
    ten = [(h["nhom"], h["ten"]) for h in r["hang_hoa"]]
    # Dịch vụ → Thuốc → Vật tư; trong nhóm theo tên.
    assert ten == [
        ("DICH_VU", "Khám phụ khoa"),
        ("DICH_VU", "Siêu âm"),
        ("THUOC", "Paracetamol"),
        ("VAT_TU", "Que thử"),
    ]
    kham, sa, para, que = r["hang_hoa"]
    assert (kham["so_luong"], kham["tien_sau_giam"]) == (2, 300_000)
    # 300.000 − 30.000 giảm = 270.000 (+ 600.000) → 870.000; SL 1 + 2.
    assert sa["so_luong"] == 3 and sa["so_dong"] == 2
    assert (sa["tien_truoc_giam"], sa["giam_gia"], sa["tien_sau_giam"]) == (
        900_000,
        30_000,
        870_000,
    )
    assert para["dvt"] == "Viên" and para["tien_sau_giam"] == 50_000
    assert que["tien_sau_giam"] == 90_000
    # Thuế để trống ở G5a (G5b điền).
    assert all(h["thue_suat"] is None and h["tien_thue"] is None for h in r["hang_hoa"])
    assert r["so_mat_hang"] == 4
    assert r["tong_truoc_giam"] == 1_340_000
    assert r["tong_giam_gia"] == 30_000
    assert r["tong_sau_giam"] == 1_310_000
    assert r["tong_truoc_giam"] - r["tong_giam_gia"] == r["tong_sau_giam"]


def test_cung_ten_khac_dvt_la_hai_mat_hang() -> None:
    r = gom_theo_hang_hoa(
        [
            {
                "dong": [
                    _d("Vitamin", 1, 10_000, dvt="Viên"),
                    _d("Vitamin", 1, 90_000, dvt="Hộp"),
                ]
            }
        ]
    )
    assert [h["dvt"] for h in r["hang_hoa"]] == ["Hộp", "Viên"]


def test_rong_va_rac_khong_nem() -> None:
    assert gom_theo_hang_hoa([])["so_mat_hang"] == 0
    r = gom_theo_hang_hoa([{"dong": None}, {}, {"dong": [{"ten": None}]}])
    assert r["so_mat_hang"] == 1 and r["tong_sau_giam"] == 0


def test_so_luong_le_va_excel_khong_thanh_cong_thuc() -> None:
    r = gom_theo_hang_hoa(
        [
            {
                "dong": [
                    _d('=HYPERLINK("x")', 0.5, 1_000),
                    _d('=HYPERLINK("x")', 0.25, 500),
                ]
            }
        ]
    )
    assert r["hang_hoa"][0]["so_luong"] == 0.75
    ws: Any = load_workbook(BytesIO(xlsx_theo_hang_hoa(r))).active
    assert [c.value for c in ws[1]][:3] == ["Nhóm", "Mã hàng", "Tên hàng hoá / dịch vụ"]
    assert ws.cell(2, 3).data_type == "s"
    assert ws.cell(2, 3).value == '=HYPERLINK("x")'
    assert ws.cell(3, 1).value == "Tổng" and ws.cell(3, 8).value == 1_500
