"""Hoá đơn công ty (G5c): luật MST chỉ nhắc + HĐ công ty trong file xuất (hàm thuần)."""

import io
import warnings
from typing import Any

import pytest
from openpyxl import load_workbook

from clinicai.core.exceptions import ValidationError
from clinicai.services.hoa_don_cong_ty_service import (
    BanGhiCongTy,
    chuan_hoa_mst,
    mst_dung_dang,
    nhac_mst,
)
from clinicai.services.xuat_hoa_don_service import DONG_DAU, dung_hoa_don, xlsx_misa

warnings.simplefilter("ignore", UserWarning)


@pytest.mark.parametrize(
    "mst",
    ["0123456789", "0123456789-001", "0123 456 789", " 0123456789 ", "0123456789-999"],
)
def test_mst_dung_dang_khong_nhac(mst: str) -> None:
    assert mst_dung_dang(mst) is True
    assert nhac_mst(mst) is None


@pytest.mark.parametrize(
    "mst",
    ["012345", "abc", "01234567890", "0123456789-01", "0123456789-0011", "0123456789-"],
)
def test_mst_sai_dang_chi_nhac(mst: str) -> None:
    assert mst_dung_dang(mst) is False
    nhac = nhac_mst(mst)
    assert nhac and "kiểm lại" in nhac
    # Nhắc trung tính: không chữ "lỗi", "sai", "không hợp lệ" (quy tắc bớt cảnh báo đỏ).
    assert not any(t in nhac.lower() for t in ("lỗi", "sai", "không hợp lệ"))


@pytest.mark.parametrize("rac", [None, 123, 4.5, [], {}, b"0123456789"])
def test_mst_rac_tra_rong_khong_nem(rac: Any) -> None:
    assert chuan_hoa_mst(rac) == ""
    assert mst_dung_dang(rac) is None
    assert nhac_mst(rac) is None


def test_mst_rong_khong_nhac() -> None:
    assert mst_dung_dang("") is None and mst_dung_dang("   ") is None
    assert nhac_mst("") is None


def test_ghi_rac_bao_loi_khong_ghi_rong() -> None:
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty="   ")
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty=None)
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty=123)
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty="A" * 401)
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty="A", dia_chi_cong_ty="x" * 501)
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty="A", ma_so_thue="1" * 21)
    with pytest.raises(ValidationError):
        BanGhiCongTy(ten_cong_ty="A", email_nhan_hd=["x"])
    # MST sai dạng thì KHÔNG ném — chỉ nhắc.
    ban = BanGhiCongTy(ten_cong_ty=" Công ty A ", ma_so_thue="12 34")
    assert ban.ten == "Công ty A" and ban.mst == "1234"


def _phieu(id_: str) -> dict[str, Any]:
    return {
        "id": id_,
        "ma": f"PT-{id_}",
        "kind": "dich_vu",
        "thu_luc": "2026-10-12T08:00:00+07:00",
        "thu_trong_ky": True,
        "trang_thai": "DA_THU",
        "hinh_thuc": "TRANSFER",
        "dong": [
            {"id": f"l-{id_}", "ten": "Khám", "so_luong": 1, "thanh_tien": 300_000}
        ],
        "hoan": [],
    }


KHACH = [
    {
        "ten": "Nguyễn Thị A",
        "ma_bn": "BN1",
        "so_dien_thoai": "0901234567",
        "phieu": [_phieu("p1"), _phieu("p2")],
    }
]
THEM_KHACH = {
    pid: {
        "dia_chi": "12 Kim Ngưu",
        "cccd": "001190000001",
        "email": "ca-nhan@example.com",
    }
    for pid in ("p1", "p2")
}
CONG_TY = {
    "p1": {
        "ten_cong_ty": "Công ty TNHH Hoa Sen",
        "dia_chi_cong_ty": "99 Trần Duy Hưng, Hà Nội",
        "ma_so_thue": "0123456789",
        "email_nhan_hd": "ketoan@hoasen.vn",
        "nguoi_mua": "",
    }
}


def test_hd_cong_ty_chi_xuat_thong_tin_cong_ty() -> None:
    hds = dung_hoa_don(KHACH, {}, THEM_KHACH, None, None, CONG_TY)
    ct, ca_nhan = hds[0], hds[1]
    assert ct["payment_cycle_id"] == "p1" and ca_nhan["payment_cycle_id"] == "p2"
    assert ct["la_hd_cong_ty"] is True
    assert (ct["ten_cong_ty"], ct["dia_chi"], ct["ma_so_thue"]) == (
        "Công ty TNHH Hoa Sen",
        "99 Trần Duy Hưng, Hà Nội",
        "0123456789",
    )
    assert ct["email"] == "ketoan@hoasen.vn"
    # Không một thông tin cá nhân nào lọt vào hoá đơn công ty.
    assert ct["so_dien_thoai"] == "" and ct["cccd"] == ""
    assert ct["nguoi_mua"] == ""
    blob = repr(ct)
    for rieng_tu in ("0901234567", "001190000001", "12 Kim Ngưu", "ca-nhan@example"):
        assert rieng_tu not in blob
    # Hoá đơn cá nhân cùng kỳ GIỮ NGUYÊN như trước G5c.
    assert not ca_nhan["la_hd_cong_ty"]
    assert (ca_nhan["dia_chi"], ca_nhan["cccd"], ca_nhan["email"]) == (
        "12 Kim Ngưu",
        "001190000001",
        "ca-nhan@example.com",
    )
    assert ca_nhan["so_dien_thoai"] == "0901234567"
    assert ca_nhan["nguoi_mua"] == "Nguyễn Thị A" and ca_nhan["ma_so_thue"] == ""


def test_hd_cong_ty_khong_doi_so_tien() -> None:
    """Chọn công ty chỉ đổi người mua — KHÔNG đổi một đồng nào (Q15)."""
    co = dung_hoa_don(KHACH, {}, THEM_KHACH, None, None, CONG_TY)
    khong = dung_hoa_don(KHACH, {}, THEM_KHACH)
    for a, b in zip(co, khong, strict=True):
        assert (a["tong"], a["dong"]) == (b["tong"], b["dong"])


def test_nguoi_mua_dien_tay_thi_ghi() -> None:
    ct = {"p1": {**CONG_TY["p1"], "nguoi_mua": "Trần Văn B"}}
    assert dung_hoa_don(KHACH, {}, THEM_KHACH, None, None, ct)[0]["nguoi_mua"] == (
        "Trần Văn B"
    )


def test_excel_cot_cde_va_hi_trong() -> None:
    hds = dung_hoa_don(KHACH, {}, THEM_KHACH, None, None, CONG_TY)
    ws: Any = load_workbook(io.BytesIO(xlsx_misa(hds))).active
    dong = {
        "ct": [ws.cell(DONG_DAU, c).value for c in range(3, 10)],
        "cn": [ws.cell(DONG_DAU + 1, c).value for c in range(3, 10)],
    }
    # C tên cty · D địa chỉ cty · E MST · F người mua trống · G email cty · H/I trống.
    assert dong["ct"] == [
        "Công ty TNHH Hoa Sen",
        "99 Trần Duy Hưng, Hà Nội",
        "0123456789",
        None,
        "ketoan@hoasen.vn",
        None,
        None,
    ]
    assert dong["cn"][0] == "Nguyễn Thị A" and dong["cn"][5] == "0901234567"
    assert dong["cn"][6] == "001190000001"


def test_cong_ty_rong_ten_khong_thanh_hd_cong_ty() -> None:
    """Bản ghi rác thiếu tên → coi như không chọn công ty (không xuất dòng trống)."""
    rac = {"p1": {"ten_cong_ty": "", "dia_chi_cong_ty": "x", "ma_so_thue": "1"}}
    hd = dung_hoa_don(KHACH, {}, THEM_KHACH, None, None, rac)[0]
    assert not hd["la_hd_cong_ty"] and hd["ten_cong_ty"] == ""
    assert hd["cccd"] == "001190000001"
